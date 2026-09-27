"""Local XLSX presentation tokens. Canonical and portable records stay unchanged."""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from urllib.parse import urlsplit

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

PERSIAN_FONT, TECHNICAL_FONT = 'Vazir', 'Poppins'
BODY_SIZE, HEADER_SIZE = 11, 11
HEADER_HEIGHT, ROW_MIN, ROW_MAX = 30, 24, 60
HEADER_FILL, HEADER_TEXT = '233E4A', 'FFFFFF'
BODY_TEXT, BAND, SEPARATOR, LINK = '243742', 'F4F7F8', 'DFE7EA', '155E91'
BADGES = {
    'بررسی نشده': ('EDF1F4', '465967'), 'علاقه‌مندم': ('E6EFFB', '234F83'),
    'می‌خوام برم': ('DDF2ED', '17604E'), 'رفتم': ('EBE5F5', '58417D'),
    'رد شد': ('F8E7E8', '8A343A'), 'یکی هستند': ('FBEED6', '785318'),
    'جدا هستند': ('DDF2ED', '17604E'), 'کامل': ('DDF2ED', '17604E'),
    'نیاز به بررسی': ('FBEED6', '785318'), 'تاریخ نامشخص': ('FBEED6', '785318'),
    'اطلاعات ناقص': ('EDF1F4', '465967'), 'تعارض داده': ('F8E7E8', '8A343A'),
    'تعارض تاریخ': ('F8E7E8', '8A343A'), 'تعارض مکان': ('F8E7E8', '8A343A'),
    'پیش رو': ('DDF2ED', '17604E'), 'پایان یافته': ('EDF1F4', '465967'),
    'نامشخص': ('EDF1F4', '465967'), 'لغو شده': ('F8E7E8', '8A343A'),
    'تکمیل ظرفیت': ('FBEED6', '785318'), 'به تعویق افتاده': ('FBEED6', '785318'),
    'موفق': ('DDF2ED', '17604E'), 'با خطا': ('FBEED6', '785318'), 'ناموفق': ('F8E7E8', '8A343A'),
}
RUN_STATUS = {'success': 'موفق', 'partial': 'با خطا', 'failed': 'ناموفق'}
CATEGORIES = {'workshop':'کارگاه', 'exhibition':'نمایشگاه', 'concert':'کنسرت',
              'festival':'جشنواره', 'conference':'همایش', 'talk':'نشست', 'meetup':'دورهمی',
              'screening':'اکران', 'market':'بازارچه', 'tour':'تور',
              'cafe':'کافه', 'gallery':'گالری', 'museum':'موزه', 'restaurant':'رستوران',
              'park':'پارک', 'bookstore':'کتاب‌فروشی'}
BADGE_HEADERS = {'تصمیم من', 'تصمیم', 'وضعیت رویداد', 'کیفیت داده', 'وضعیت', 'ظرفیت / وضعیت ثبت‌نام'}


@dataclass(frozen=True)
class ColumnStyle:
    minimum: float
    preferred: float
    maximum: float
    wrap: bool = False
    centered: bool = False


COMPACT = ColumnStyle(12, 15, 18, centered=True)
DATE = ColumnStyle(17, 19, 22, centered=True)
TIME = ColumnStyle(10, 12, 14, centered=True)
TITLE = ColumnStyle(24, 30, 40, wrap=True)
TEXT = ColumnStyle(24, 32, 44, wrap=True)
DESCRIPTION_TEXT = ColumnStyle(36, 44, 56, wrap=True)
SHORT_TEXT = ColumnStyle(16, 20, 26, wrap=True)
LINK_COLUMN = ColumnStyle(14, 16, 20)
DEFAULT = ColumnStyle(14, 18, 24)
GUIDE_TEXT = ColumnStyle(64, 74, 84, wrap=True)


def font_name(value: object) -> str:
    return PERSIAN_FONT if re.search(r'[\u0600-\u06ff]', str(value or '')) else TECHNICAL_FONT


def column_style(header: str) -> ColumnStyle:
    if header == 'راهنمای مرور فرصت‌ها':
        return GUIDE_TEXT
    if header in BADGE_HEADERS or header in {'دسته‌بندی','شهر','روز','فرمت','تعداد منابع','مدت'}:
        return COMPACT
    if header == 'توضیحات / معرفی':
        return DESCRIPTION_TEXT
    if 'تاریخ' in header:
        return DATE
    if 'ساعت' in header or header == 'زمان اجرا':
        return TIME
    if header in {'منطقه / محله','برگزارکننده'}:
        return SHORT_TEXT
    if header in {'عنوان','نام مکان','رویداد A','رویداد B','مکان'}:
        return TITLE
    if header in {'آدرس','خلاصه','یادداشت من','دلیل پیشنهاد','خطاها','زمان‌بندی اعلام‌شده'}:
        return TEXT
    if 'منبع' in header or 'منابع' in header or header == 'ثبت‌نام / خرید':
        return LINK_COLUMN
    return DEFAULT


def link_label(target: str, header: str) -> str:
    if header == 'ثبت‌نام / خرید':
        return header
    host = urlsplit(target).hostname or ''
    for domain, label in (('davvvat.ir','دعوت'),('vadoostan.ir','ودوستان'),('jabama.events','جاباما'),('instagram.com','Instagram')):
        if host == domain or host.endswith('.' + domain):
            return label
    return 'مشاهده منبع'


def apply_design(sheet, headers: tuple, visible: int) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'A2'
    sheet.row_dimensions[1].height = HEADER_HEIGHT
    edge = Border(bottom=Side(style='hair', color=SEPARATOR))
    for column, header in enumerate(headers, 1):
        spec = column_style(header)
        letter = get_column_letter(column)
        dimension = sheet.column_dimensions[letter]
        dimension.hidden = column > visible
        # URL targets never participate in sizing. Only bounded display text does.
        lengths = [max(map(len,str(sheet.cell(r,column).value or '').splitlines()),default=0)
                   for r in range(1,min(sheet.max_row,101)+1)]
        content = max(lengths, default=0) * 0.85 + 2
        dimension.width = max(spec.minimum,min(spec.maximum,max(spec.preferred,content)))
        for row in range(1,sheet.max_row+1):
            cell = sheet.cell(row,column)
            header_cell = row == 1
            name = PERSIAN_FONT if header_cell and column <= visible else font_name(cell.value)
            fill = HEADER_FILL if header_cell else (BAND if row % 2 == 0 else 'FFFFFF')
            color = HEADER_TEXT if header_cell else BODY_TEXT
            badge = header in BADGE_HEADERS and row > 1
            if badge and cell.value in BADGES:
                fill, color = BADGES[cell.value]
            if cell.hyperlink and not header_cell:
                color = LINK
            cell.font = Font(name=name, size=HEADER_SIZE if header_cell else BODY_SIZE,
                             bold=header_cell or badge, color=color, underline='single' if cell.hyperlink else None)
            cell.fill = PatternFill('solid',fgColor=fill)
            cell.border = edge
            cell.alignment = Alignment(vertical='center', horizontal='center' if header_cell or spec.centered else
                                       ('right' if name==PERSIAN_FONT else 'left'),
                                       wrap_text=header_cell or spec.wrap, readingOrder=2 if name==PERSIAN_FONT else 1)
        if header in {'تصمیم من','تصمیم'}:
            for value,(fill,color) in BADGES.items():
                if value not in {'بررسی نشده','علاقه‌مندم','می‌خوام برم','رفتم','رد شد','یکی هستند','جدا هستند'}:
                    continue
                sheet.conditional_formatting.add(f'{letter}2:{letter}{max(1000,sheet.max_row+100)}',
                    CellIsRule(operator='equal',formula=['"'+value+'"'],
                               fill=PatternFill('solid',fgColor=fill),font=Font(name=PERSIAN_FONT,bold=True,color=color)))
    for row in range(2,sheet.max_row+1):
        lines = 1
        for column,header in enumerate(headers[:visible],1):
            if column_style(header).wrap:
                value = str(sheet.cell(row,column).value or '')
                width = sheet.column_dimensions[get_column_letter(column)].width
                lines = max(lines,sum(max(1,math.ceil(len(part)/(width*1.1))) for part in value.splitlines()))
        sheet.row_dimensions[row].height = min(ROW_MAX, max(ROW_MIN, lines*18+6))
