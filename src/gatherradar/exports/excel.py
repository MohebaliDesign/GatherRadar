"""Persian local review workspace and narrowly scoped human-state import."""
from __future__ import annotations

import os
from collections import Counter
from dataclasses import dataclass, fields
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from ..domain import Event, PlaceCandidate
from ..review.models import CanonicalRepository, ReviewError
from ..review.records import event_from_record
from ..review.relevance import relevant_document
from ..review.schema import (DECISIONS, DUPLICATE_DECISIONS, DUPLICATE_HEADERS,
    EVENT_HEADERS, EVENT_VISIBLE, PLACE_HEADERS, PLACE_VISIBLE, GUIDE, GUIDE_ROWS)
from ..review.serialization import event_row, place_row, snapshot_title, url, jalali, persian
from .excel_design import LINK_COLUMN, apply_design, column_style, link_label, CATEGORIES, RUN_STATUS

# 3: source schedule/area/duration/availability/description/organizer columns.
# Version 2 workbooks differ only in new snapshot columns; they are upgraded on
# the next export without rewriting historical snapshots. Older/newer refuse.
WORKBOOK_VERSION = 3
READABLE_VERSIONS = (2, 3)
XLSX_CELL_LIMIT = 32767
DESCRIPTION = 'توضیحات / معرفی'
TRUNCATED = '\n… [متن کامل در review.json و events.csv]'
RUNS, DUPLICATES = 'تاریخچه اجراها', 'بررسی تکراری‌ها'
TOKEN = 'review_token'
DUPLICATE_HEADERS = DUPLICATE_HEADERS + ('source_urls_a','source_urls_b')
INDEX_HEADERS = ('شناسه اجرا', 'تاریخ اجرا', 'وضعیت', 'رویدادهای نمایش داده', 'مکان‌ها', 'تکراری احتمالی',
                 'بازه (روز)', 'خطاها', 'تب رویدادها', 'تب مکان‌ها', 'زمان اجرا',
                 'تعداد منابع', 'کل رویدادها', 'شروع اجرا (ISO)')
LOCAL_EVENT_VISIBLE = EVENT_VISIBLE
NEXT_HEADERS = ('جلسه بعدی', 'شروع جلسه بعدی', 'پایان جلسه بعدی', 'جلسه‌های باقی‌مانده')


def open_workspace(path: Path) -> Workbook:
    try:
        if path.suffix.lower() != '.xlsx' or path.stat().st_size > 50_000_000:
            raise ValueError
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 200_000_000:
                raise ValueError
        workbook = load_workbook(path, data_only=False, keep_links=False)
        if (any(s.max_row > 100000 or s.max_column > 256 for s in workbook)
                or '_meta' not in workbook or workbook['_meta']['A1'].value != 'schema_version'
                or workbook['_meta']['B1'].value not in READABLE_VERSIONS
                or any(name not in workbook for name in (GUIDE, RUNS, DUPLICATES))
                or tuple(c.value for c in workbook[RUNS][1]) != INDEX_HEADERS
                or tuple(c.value for c in workbook[DUPLICATES][1]) != DUPLICATE_HEADERS + (TOKEN,)):
            workbook.close()
            raise ValueError
        return workbook
    except Exception:
        raise ReviewError('Existing XLSX is unreadable or has an incompatible schema. It was not overwritten; restore it or choose a separate --workbook recovery path.') from None


# Oversized snapshot values are stored losslessly in this hidden sheet, in chunks
# below the cell limit, keyed by snapshot sheet, record ID and field. The snapshot
# cell holds a reference marker (technical fields) or a readable prefix with a
# visible marker. Historical overflow rows are appended to, never rewritten.
OVERFLOW = '_overflow'
OVERFLOW_HEADERS = ('sheet', 'record_id', 'field', 'sha256', 'length', 'chunk', 'chunks', 'text')
OVERFLOW_CHUNK = 32_000
OVERFLOW_REFERENCE = '[[gatherradar-overflow sha256={digest} length={length} chunks={chunks}]]'
OVERFLOW_VISIBLE = '\n… [متن کامل در برگه پنهان _overflow و در review.json]'
# Owner-entered values round-trip through import; shortening them would lose data.
REVIEWER_OWNED = frozenset({'تصمیم من', 'یادداشت من', 'تصمیم'})


class CellLimitError(ReviewError):
    def __init__(self, value_length: int, *, sheet: str | None = None, record_id: str | None = None,
                 field: str | None = None) -> None:
        self.value_length, self.sheet, self.record_id, self.field = value_length, sheet, record_id, field
        where = ', '.join(f'{k}={v}' for k, v in (('sheet', sheet), ('record_id', record_id),
                                                  ('field', field)) if v)
        super().__init__(f'A review cell exceeds the XLSX text limit ({value_length} > {XLSX_CELL_LIMIT}'
                         + (f'; {where}' if where else '') + '); use JSON/CSV and inspect the source.')


def literal(cell, value: object) -> None:
    if isinstance(value, str):
        # openpyxl would silently truncate long strings; fail instead of losing facts.
        if len(value) > XLSX_CELL_LIMIT:
            raise CellLimitError(len(value))
        cell.value = value
        cell.data_type = 's'
    else:
        cell.value = value


def fit_rows(sheet_name: str, headers: tuple, rows: list[list], key_header: str,
             visible: int) -> tuple[list[list], list[tuple]]:
    """Cells within the XLSX limit plus lossless overflow rows for every longer value."""
    key = headers.index(key_header)
    fitted, overflow = [], []
    for row in rows:
        row = list(row)
        for column, value in enumerate(row):
            if not isinstance(value, str) or len(value) <= XLSX_CELL_LIMIT:
                continue
            header = headers[column]
            if header in REVIEWER_OWNED:
                raise CellLimitError(len(value), sheet=sheet_name, record_id=str(row[key]), field=header)
            digest = sha256(value.encode('utf-8')).hexdigest()
            chunks = -(-len(value) // OVERFLOW_CHUNK)
            overflow.extend((sheet_name, row[key], header, digest, len(value), index + 1, chunks,
                             value[index * OVERFLOW_CHUNK:(index + 1) * OVERFLOW_CHUNK])
                            for index in range(chunks))
            if column < visible:
                marker = TRUNCATED if header == DESCRIPTION else OVERFLOW_VISIBLE
                row[column] = value[:XLSX_CELL_LIMIT - len(marker)] + marker
            else:
                row[column] = OVERFLOW_REFERENCE.format(digest=digest, length=len(value), chunks=chunks)
        fitted.append(row)
    return fitted, overflow


def store_overflow(workbook: Workbook, rows: list[tuple]) -> None:
    if not rows:
        return
    if OVERFLOW not in workbook:
        sheet = workbook.create_sheet(OVERFLOW)
        sheet.sheet_state = 'hidden'
        for column, header in enumerate(OVERFLOW_HEADERS, 1):
            literal(sheet.cell(1, column), header)
    sheet = workbook[OVERFLOW]
    if tuple(c.value for c in sheet[1]) != OVERFLOW_HEADERS:
        raise ReviewError('The hidden _overflow sheet has an unexpected layout; it was not modified.')
    start = sheet.max_row + 1
    for offset, values in enumerate(rows):
        for column, value in enumerate(values, 1):
            literal(sheet.cell(start + offset, column), value)


def overflow_value(workbook: Workbook, sheet_name: str, record_id: str, field: str) -> str:
    """Reassemble one stored value; refuse incomplete or altered chunks."""
    if OVERFLOW not in workbook:
        raise ReviewError('This workbook has no stored overflow values.')
    parts = [r for r in workbook[OVERFLOW].iter_rows(min_row=2, values_only=True)
             if r[:3] == (sheet_name, record_id, field)]
    if not parts:
        raise ReviewError(f'No overflow value for sheet={sheet_name}, record_id={record_id}, field={field}.')
    digest, length, chunks = parts[0][3], parts[0][4], parts[0][6]
    ordered = sorted(parts, key=lambda r: r[5])
    value = ''.join(r[7] or '' for r in ordered)
    if ([r[5] for r in ordered] != list(range(1, chunks + 1)) or any(r[3] != digest for r in ordered)
            or len(value) != length or sha256(value.encode('utf-8')).hexdigest() != digest):
        raise ReviewError(f'Overflow value for {record_id}/{field} is incomplete or altered.')
    return value


def table(sheet, headers: tuple, rows: list[list], visible: int, decisions: tuple = ()) -> None:
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(1, len(rows)+1)}'
    key = next((h for h in ('event_id', 'candidate_id', 'pair_key', 'شناسه اجرا') if h in headers), None)
    for r, values in enumerate([list(headers), *rows], 1):
        for c, value in enumerate(values, 1):
            cell = sheet.cell(r, c)
            try:
                literal(cell, value)
            except CellLimitError as exc:
                record = str(values[headers.index(key)]) if key and r > 1 else None
                raise CellLimitError(exc.value_length, sheet=sheet.title, record_id=record,
                                     field=str(headers[c-1])) from None
            # Only link columns become hyperlinks; a description that happens to
            # start with a URL keeps its full text.
            if r > 1 and c <= visible and isinstance(value, str) and (
                    column_style(headers[c-1]) is LINK_COLUMN or headers[c-1].startswith('پیوند منبع')):
                try:
                    if value.startswith(('https://', 'http://')):
                        targets = value.splitlines()
                        if not all(url(target) for target in targets):
                            raise ValueError
                        cell.hyperlink = targets[0]
                        label = link_label(targets[0], headers[c-1])
                        literal(cell, label + (f' ({persian(len(targets))} منبع)' if len(targets)>1 else ''))
                except ValueError:
                    pass
    apply_design(sheet, headers, visible)
    if decisions:
        validation = DataValidation(type='list', formula1='"' + ','.join(decisions) + '"', allow_blank=False)
        validation.errorTitle = 'مقدار نامعتبر'
        validation.error = 'یکی از گزینه‌های فهرست را انتخاب کنید.'
        validation.showErrorMessage = True
        validation.errorStyle = 'stop'
        sheet.add_data_validation(validation)
        validation.add(f'A2:A{max(1000, len(rows)+100)}')


def _name(kind: str, run: dict, workbook: Workbook) -> str:
    name = snapshot_title(kind, datetime.fromisoformat(run['started_at']), run['review_timezone']).replace('_',' ')
    if len(name) > 31 or name in workbook:
        suffix = sha256(run['run_id'].encode()).hexdigest()[:7]
        name = name[:23] + '_' + suffix
    if name in workbook:
        raise ReviewError('Snapshot name collision; no historical sheet was overwritten.')
    return name


def _duplicate_row(item: dict, token: str) -> list:
    return [item['review_decision'], item['title_a'], item['jalali_date_a'], '\n'.join(item['source_urls_a']),
            item['title_b'], item['jalali_date_b'], '\n'.join(item['source_urls_b']), item['reason'],
            item['review_notes'], item['pair_key'], item['event_id_a'], item['event_id_b'],
            item['first_seen_run'], item['last_seen_run'], '\n'.join(item['reason_codes']),
            '\n'.join(item['source_urls_a']),'\n'.join(item['source_urls_b']), token]


class ExcelReviewExporter:
    def __init__(self, path: str | Path, repository: CanonicalRepository) -> None:
        self.path, self.repository = Path(path), repository
        if self.path.suffix.lower() != '.xlsx':
            raise ReviewError('Local review requires a .xlsx workbook path.')

    def export(self, document: dict) -> tuple[Path, ...]:
        document = relevant_document(document)
        workbook = open_workspace(self.path) if self.path.exists() else self._new()
        temporary = self.path.with_name(self.path.stem + '.' + uuid4().hex + '.tmp.xlsx')
        try:
            run = document['run']
            run_id = run['run_id']
            index = workbook[RUNS]
            known = [list(row) for row in index.iter_rows(min_row=2, values_only=True)]
            ids = [r[0] for r in known]
            if len(set(ids)) != len(ids):
                raise ReviewError('Duplicate run IDs in XLSX index; restore the index before export.')
            for row in known:
                if any(name and name not in workbook for name in row[8:10]):
                    raise ReviewError('Historical XLSX snapshot is missing; restore it or choose a recovery workbook.')
            previous = next((row for row in known if row[0] == run_id), None)
            from ..review.records import dumps
            # Every factual change must reach a fresh projection, including price,
            # duration and description. Reviewer-owned edits do not create tabs.
            projection_fields = tuple(f.name for f in fields(Event)
                if f.name not in {'review_status', 'reviewer_notes'}) + (
                'schedule_reason', 'next_session_date', 'next_session_start_time',
                'next_session_end_time', 'remaining_session_count', 'readiness_label', 'readiness_explanation')
            projection = sha256(dumps([(kind, [
                {key: r.get(key) for key in projection_fields} for r in document.get(kind, [])])
                for kind in ('events', 'schedule_review')]).encode()).hexdigest()
            saved_projection = next((r[2] for r in workbook['_meta'].iter_rows(values_only=True)
                if len(r) >= 3 and r[0] == 'schedule_projection' and r[1] == run_id), None)
            superseded = None
            if previous and previous[8]:
                sheet = workbook[previous[8]]
                headers = [c.value for c in sheet[1]]
                if headers.count('event_id') != 1:
                    raise ReviewError('Historical Event identity header is invalid; snapshot was preserved.')
                column = headers.index('event_id')
                old_ids = [row[column] for row in sheet.iter_rows(min_row=2, values_only=True)]
                if (set(old_ids) != {e['event_id'] for e in document['events']}
                        or saved_projection != projection):
                    superseded = previous
                    known.remove(previous)
            if previous is None or superseded:
                naming_run = dict(run)
                if superseded:
                    # A new projection of the same run gets its own sheet; the
                    # old sheet's cells, IDs, review values and layout stay intact.
                    naming_run['run_id'] += '-projection-' + str(len(workbook.sheetnames))
                    meta = workbook['_meta']
                    meta.append(('superseded_event_sheet', superseded[8]))
                    prior_schedules = [r[2] for r in meta.iter_rows(values_only=True)
                        if len(r) >= 3 and r[0] == 'schedule_review_sheet' and r[1] == run_id]
                    for old in prior_schedules:
                        meta.append(('superseded_event_sheet', old))
                event_name = _name('رویدادها', naming_run, workbook)
                event_rows = []
                # Evidence first, then public reference URLs stated in the evidence
                # (clickable for the owner; their unseen page content is not claimed).
                all_links = [(lambda seen: seen + [u for u in dict.fromkeys(item.get('reference_urls') or ())
                                                   if u not in seen])(sorted(set(item['evidence_urls'] + [item['canonical_source_url']])))
                             for item in document['events']]
                link_count = max((len(links) for links in all_links), default=0)
                # OOXML allows one hyperlink per cell. Extra supporting-source columns
                # keep every URL clickable without formulas or invented links.
                link_count = link_count if link_count > 1 else 0
                link_headers = tuple(f'پیوند منبع {persian(i+1)}' for i in range(link_count))
                for item, links in zip(document['events'], all_links):
                    token = self.repository.review_token('event', item['event_id'], run_id)
                    row = event_row(event_from_record(item), run_id, (item['review_decision'], item['review_notes']))
                    values = dict(zip(EVENT_HEADERS,row))
                    values['آمادگی داده'] = item.get('readiness_label', values['آمادگی داده'])
                    values['دلیل بررسی'] = item.get('readiness_explanation', values['دلیل بررسی'])
                    values['data_readiness'] = item.get('data_readiness', values['data_readiness'])
                    from ..review.records import dumps
                    values['readiness_reasons'] = dumps(item.get('readiness_reasons', []))
                    values['quality_indicators'] = dumps(item.get('quality_indicators', []))
                    # Unmapped explicit source category is shown as its exact wording.
                    values['دسته‌بندی'] = (CATEGORIES.get(item['category'],item['category'])
                                           or item.get('source_category_text') or '')
                    event_rows.append([values[h] for h in LOCAL_EVENT_VISIBLE] +
                        [jalali(datetime.fromisoformat(item['next_session_date']).date()) if item.get('next_session_date') else '',
                         persian((item.get('next_session_start_time') or '')[:5]),
                         persian((item.get('next_session_end_time') or '')[:5]), item.get('remaining_session_count'),
                         item.get('location_context') or ''] +
                        (links + [''] * (link_count-len(links)) if link_count else []) +
                        [values[h] for h in EVENT_HEADERS[len(EVENT_VISIBLE):]] +
                        [item['category'],item['event_format'],item['status'],item['registration_url'],token])
                headers = (LOCAL_EVENT_VISIBLE + NEXT_HEADERS + ('موقعیت تکمیلی',) + link_headers + EVENT_HEADERS[len(EVENT_VISIBLE):]
                           + ('category','event_format','status','registration_url',TOKEN))
                visible = len(LOCAL_EVENT_VISIBLE) + len(NEXT_HEADERS) + 1 + link_count
                # Over-limit values (e.g. a long session list) are kept whole in _overflow;
                # a description keeps its explicit marker; SQLite/JSON/CSV are unchanged.
                event_rows, overflow = fit_rows(event_name, headers, event_rows, 'event_id', visible)
                table(workbook.create_sheet(event_name, 3), headers, event_rows, visible, DECISIONS)
                store_overflow(workbook, overflow)
                if document.get('schedule_review'):
                    review_name = _name('زمان‌بندی', naming_run, workbook)
                    rows = []
                    for item in document['schedule_review']:
                        row = event_row(event_from_record(item), run_id,
                                        (item['review_decision'], item['review_notes']))
                        row[EVENT_HEADERS.index('دسته‌بندی')] = CATEGORIES.get(item['category'], item['category']) or item.get('source_category_text') or ''
                        for heading, field in (('آمادگی داده', 'readiness_label'),
                                ('دلیل بررسی', 'readiness_explanation'), ('data_readiness', 'data_readiness')):
                            row[EVENT_HEADERS.index(heading)] = item.get(field, row[EVENT_HEADERS.index(heading)])
                        for heading in ('readiness_reasons', 'quality_indicators'):
                            row[EVENT_HEADERS.index(heading)] = dumps(item.get(heading, []))
                        rows.append(row + [item['schedule_reason'], item.get('location_context') or '',
                            self.repository.review_token('event', item['event_id'], run_id)])
                    review_headers = EVENT_HEADERS + ('schedule_reason', 'location_context', TOKEN)
                    rows, overflow = fit_rows(review_name, review_headers, rows, 'event_id', len(EVENT_VISIBLE))
                    table(workbook.create_sheet(review_name, 4), review_headers, rows, len(EVENT_VISIBLE), DECISIONS)
                    store_overflow(workbook, overflow)
                    workbook['_meta'].append(('schedule_review_sheet', run_id, review_name))
                meta = workbook['_meta']
                for row in meta.iter_rows():
                    if len(row) >= 3 and row[0].value == 'schedule_projection' and row[1].value == run_id:
                        row[2].value = projection
                        break
                else:
                    meta.append(('schedule_projection', run_id, projection))
                place_name = superseded[9] if superseded else ''
                if document['places'] and not superseded:
                    place_name = _name('مکان‌ها', run, workbook)
                    rows = []
                    for item in document['places']:
                        place = PlaceCandidate(**{f.name: item[f.name] for f in fields(PlaceCandidate)})
                        row = place_row(place, run_id, item['source_id'] or '')
                        row[0], row[9] = item['review_decision'], item['review_notes']
                        row[2] = CATEGORIES.get(row[2],row[2])
                        rows.append(row + [self.repository.review_token('place', place.candidate_id, run_id)])
                    rows, overflow = fit_rows(place_name, PLACE_HEADERS + (TOKEN,), rows, 'candidate_id', len(PLACE_VISIBLE))
                    table(workbook.create_sheet(place_name, 4), PLACE_HEADERS + (TOKEN,), rows, len(PLACE_VISIBLE), DECISIONS)
                    store_overflow(workbook, overflow)
                stamp = datetime.fromisoformat(run['started_at'])
                from zoneinfo import ZoneInfo
                local = stamp.astimezone(ZoneInfo(run['review_timezone']))
                known.append([run_id, jalali(local.date()), RUN_STATUS.get(run['status'],run['status']),
                    run['event_count'], run['place_count'], run['duplicate_count'], run['days'],
                    '\n'.join(run['source_failures']), event_name, place_name, persian(local.strftime('%H:%M')),
                    run['source_count'],run['event_count']+run['filtered'],run['started_at']])
            # Frozen replays can share an exact reference with their original
            # run. Prefer the exported snapshot on a tie, retaining both sheets.
            known.sort(key=lambda row: (datetime.fromisoformat(row[13]), row[0] == run_id), reverse=True)
            # Permanent navigation/queue are rebuildable; historical snapshots are untouched.
            workbook.remove(index)
            table(workbook.create_sheet(RUNS, 1), INDEX_HEADERS, known, len(INDEX_HEADERS)-1)
            for row in workbook[RUNS].iter_rows(min_row=2):
                for cell in row[8:10]:
                    if cell.value:
                        cell.hyperlink = "#'" + cell.value.replace("'", "''") + "'!A1"
            workbook.remove(workbook[DUPLICATES])
            rows = [_duplicate_row(item, self.repository.review_token('duplicate', item['pair_key'], item['last_seen_run']))
                    for item in self.repository.duplicate_queue()]
            table(workbook.create_sheet(DUPLICATES, 2), DUPLICATE_HEADERS + (TOKEN,), rows, 9, DUPLICATE_DECISIONS)
            meta = workbook['_meta']
            for row, values in enumerate((('schema_version', WORKBOOK_VERSION), ('latest_generated_run', run_id),
                                           ('generated_timestamp', run['finished_at'])), 1):
                for column, value in enumerate(values, 1):
                    literal(meta.cell(row, column), value)
            apply_design(meta, ('key','value'),0)
            # Keep old snapshot content intact; navigation follows chronological runs,
            # even when an older persisted run is explicitly re-exported.
            snapshots = [name for row in known for name in row[8:10] if name]
            newest = [name for name in known[0][8:10] if name] if known else []
            if known:
                superseded_names = {r[1] for r in meta.iter_rows(values_only=True)
                                    if len(r) >= 2 and r[0] == 'superseded_event_sheet'}
                newest[1:1] = [r[2] for r in meta.iter_rows(values_only=True)
                    if len(r) >= 3 and r[0] == 'schedule_review_sheet' and r[1] == known[0][0]
                    and r[2] not in superseded_names]
            ordered = newest + [RUNS,DUPLICATES] + [name for name in snapshots if name not in newest] + [GUIDE,'_meta']
            ordered += [name for name in workbook.sheetnames if name not in ordered]
            for position,name in enumerate(ordered):
                workbook.move_sheet(name, offset=position-workbook.sheetnames.index(name))
            workbook.active = workbook.sheetnames.index(newest[0] if newest else RUNS)
            self.path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(temporary)
            with temporary.open('r+b') as handle:
                os.fsync(handle.fileno())
            verified = open_workspace(temporary)
            try:
                if verified.sheetnames != workbook.sheetnames:
                    raise ReviewError('Temporary XLSX verification failed.')
                for sheet in workbook:
                    if list(sheet.values) != list(verified[sheet.title].values):
                        # OOXML represents empty literal strings as empty cells.
                        normalized = lambda rows: [[None if v == '' else v for v in r] for r in rows]
                        if normalized(sheet.values) != normalized(verified[sheet.title].values):
                            raise ReviewError('Temporary XLSX cell verification failed.')
            finally:
                verified.close()
            temporary.replace(self.path)
        except PermissionError:
            raise ReviewError('XLSX replacement was blocked, possibly because the workbook is open. Close it and run export --run latest; the old file and SQLite state are preserved.') from None
        finally:
            workbook.close()
            temporary.unlink(missing_ok=True)
        return (self.path,)

    @staticmethod
    def _new() -> Workbook:
        workbook = Workbook()
        workbook.remove(workbook.active)
        guide = [list(row) for row in GUIDE_ROWS[1:]]
        guide[1] = ['بازخوانی داده محلی', 'python -m gatherradar refresh --all-enabled --stored']
        guide += [['ذخیره تصمیم‌ها', 'پیش از اجرای بعد فایل را ذخیره و ببندید. فقط تصمیم و یادداشت وارد SQLite می‌شود.'],
                  ['مکان‌ها', 'تصمیم مکان فقط به همین اجرا تعلق دارد؛ حذف تکرار مکان پیاده‌سازی نشده است.'],
                  ['پیوندها', 'منبع اصلی و ثبت‌نام کلیک‌پذیرند؛ منابع متعدد در ستون‌های پیوند منبع نیز نمایش داده می‌شوند.'],
                  ['درباره', 'GatherRadar فرصت‌های عمومی منابع منتخب را برای مرور و تصمیم شما گردآوری می‌کند.'],
                  ['آخرین رویدادها', 'اولین تب، جدیدترین تصویر رویدادهاست. تاریخچه اجراها پیوند تصاویر قدیمی را نگه می‌دارد.'],
                  ['اصلاح داده', 'تغییر عنوان، زمان، مکان و سایر واقعیت‌ها فقط در همین تصویر دیده می‌شود؛ به داده اصلی وارد نمی‌شود.'],
                  ['خروجی‌ها', 'CSV و JSON برای جابه‌جایی داده و بسته Gemini برای ساخت دستی شیت در پوشه exports قرار دارند.'],
                  ['قلم‌ها', 'برای بهترین نمایش Vazir و Poppins را نصب کنید؛ قلم‌ها همراه فایل نیستند و جایگزین به برنامه نمایش‌دهنده بستگی دارد.'],
                  ['توضیحات بلند', 'متن کامل در سلول است؛ ارتفاع ردیف محدود است و برای خواندن کامل، سلول را باز کنید. فقط متن بیش از سقف Excel علامت‌گذاری و کامل در JSON/CSV نگه داشته می‌شود.'],
                  ['داده‌های بسیار بلند', 'مقدار بیش از سقف یک سلول Excel (مثلاً فهرست ده‌ها جلسه) بدون حذف در برگه پنهان _overflow تکه‌تکه نگه داشته می‌شود و سلول اصلی به آن ارجاع می‌دهد.']]
        table(workbook.create_sheet(GUIDE), tuple(GUIDE_ROWS[0]), guide, 2)
        table(workbook.create_sheet(RUNS), INDEX_HEADERS, [], len(INDEX_HEADERS))
        table(workbook.create_sheet(DUPLICATES), DUPLICATE_HEADERS + (TOKEN,), [], 9, DUPLICATE_DECISIONS)
        meta = workbook.create_sheet('_meta')
        meta.sheet_state = 'hidden'
        for row in (('schema_version', WORKBOOK_VERSION), ('latest_generated_run', ''),
                    ('generated_timestamp', '')):
            meta.append(row)
        return workbook


@dataclass(frozen=True)
class ImportSummary:
    imported: int = 0
    skipped: int = 0
    malformed: int = 0
    conflicts: int = 0


class ExcelReviewImporter:
    def __init__(self, path: str | Path, repository: CanonicalRepository) -> None:
        self.path, self.repository = Path(path), repository

    def import_reviews(self) -> ImportSummary:
        if not self.path.exists():
            return ImportSummary()
        workbook = open_workspace(self.path)
        counts = Counter()
        try:
            run_order = {run['run_id']: index for index, run in enumerate(self.repository.list_runs())}
            superseded = {row[1] for row in workbook['_meta'].iter_rows(values_only=True)
                          if len(row) >= 2 and row[0] == 'superseded_event_sheet'}
            candidates = []
            for sheet in workbook:
                if sheet.title in (GUIDE, RUNS, '_meta', OVERFLOW) or sheet.title in superseded:
                    continue
                headers = [cell.value for cell in sheet[1]]
                kind = 'duplicate' if sheet.title == DUPLICATES else 'event' if 'event_id' in headers else 'place' if 'candidate_id' in headers else None
                if kind is None:
                    continue  # Owner-added sheets are never ingested.
                key_header = {'event': 'event_id', 'place': 'candidate_id', 'duplicate': 'pair_key'}[kind]
                run_header = 'last_seen_run' if kind == 'duplicate' else 'run_id'
                decision_header = 'تصمیم' if kind == 'duplicate' else 'تصمیم من'
                needed = (key_header, run_header, decision_header, 'یادداشت من', TOKEN)
                if any(headers.count(name) != 1 for name in needed):
                    raise ReviewError('XLSX review/identity headers are missing or duplicated; restore them before importing.')
                for cells in sheet.iter_rows(min_row=2):
                    if not any(c.value is not None for c in cells):
                        continue
                    values = [cells[headers.index(name)] for name in needed]
                    key, run_id, decision, notes, token = [c.value for c in values]
                    if (any(c.data_type == 'f' for c in values) or not all(isinstance(v, str) for v in (key, run_id, token))
                            or run_id not in run_order):
                        counts['malformed'] += 1
                        continue
                    candidates.append((run_order[run_id], kind, key, run_id, decision, notes if notes is not None else '', token))
            # Only newest snapshot containing each Event is eligible; sorting tabs/rows
            # cannot revive an obsolete decision from historical snapshots.
            newest = {}
            occurrences = Counter((kind, key, run_id) for _, kind, key, run_id, *_ in candidates)
            for row in candidates:
                _, kind, key, run_id, *_ = row
                identity = (kind, key, run_id if kind == 'place' else '')
                if identity not in newest or row[0] > newest[identity][0]:
                    newest[identity] = row
            for row in candidates:
                _, kind, key, run_id, decision, notes, token = row
                identity = (kind, key, run_id if kind == 'place' else '')
                if occurrences[(kind, key, run_id)] != 1:
                    counts['malformed'] += 1
                elif row != newest[identity]:
                    counts['skipped'] += 1
                else:
                    outcome = self.repository.import_review(token, kind, key, run_id, decision, notes)
                    counts['conflicts' if outcome == 'conflict' else outcome] += 1
        finally:
            workbook.close()
        return ImportSummary(**{key: counts[key] for key in ('imported', 'skipped', 'malformed', 'conflicts')})
