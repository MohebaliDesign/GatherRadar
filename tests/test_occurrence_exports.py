import csv
import json
from dataclasses import replace
from datetime import date, time
from decimal import Decimal

from openpyxl import load_workbook

from gatherradar.domain.occurrence import EventOccurrence, ScheduleKind
from gatherradar.domain.temporal import DatePrecision
from gatherradar.exports.excel import ExcelReviewExporter, ExcelReviewImporter
from gatherradar.exports.gemini import GeminiBundleExporter
from gatherradar.exports.portable import CsvReviewExporter, JsonReviewExporter
from gatherradar.review.records import event_from_record, event_record, upgrade_event_record
from gatherradar.review.schema import DECISIONS
from local_review_fakes import LocalCase


class OccurrenceOutputTests(LocalCase):
    def setUp(self):
        super().setUp()
        self.occurrences = (
            EventOccurrence("session:a", date(2026, 9, 28), "Asia/Tehran",
                start_time=time(19), source_schedule_text="۶ مهر ساعت ۱۹",
                price_amount=Decimal("150000"), currency="TOMAN", price_text="۱۵۰ هزار تومان",
                source_session_id="a", source_item_ids=("raw:a",),
                evidence_urls=("https://example.test/event",), evidence_fragment_ids=("fragment:a",)),
            EventOccurrence("session:b", date(2026, 10, 2), "Asia/Tehran",
                start_time=time(20), availability_text="تکمیل ظرفیت", source_session_id="b"),
        )
        self.event = replace(self.result.events[0], occurrences=self.occurrences,
                             start_date=None, start_time=None, starts_at=None,
                             date_precision=DatePrecision.UNKNOWN, schedule_kind=ScheduleKind.DISCRETE)
        self.result = replace(self.result, events=(self.event,), possible_duplicates=())

    def test_typed_roundtrip_preserves_sessions_and_null_parent(self):
        result = event_from_record(event_record(self.event))
        self.assertEqual(result, self.event)
        self.assertIsNone(result.start_date)
        self.assertIsNone(result.occurrences[1].price_amount)

    def test_record_migration_is_idempotent_and_does_not_invent_dates(self):
        old = event_record(self.event)
        for name in ("occurrences", "schedule_kind", "recurrence_rule"):
            old.pop(name)
        once = upgrade_event_record(old)
        self.assertEqual(upgrade_event_record(once), once)
        self.assertEqual(once["occurrences"], [])
        self.assertIsNone(once["start_date"])

    def test_persistence_portable_and_gemini_agree(self):
        self.persist()
        document = self.repo.load()
        exports = self.root / "exports"
        CsvReviewExporter(exports).export(document)
        JsonReviewExporter(exports).export(document)
        GeminiBundleExporter(exports / "gemini").export(document)
        actual = json.loads((exports / "latest/review.json").read_text(encoding="utf-8"))
        with (exports / "latest/events.csv").open(encoding="utf-8-sig", newline="") as handle:
            csv_row = next(csv.DictReader(handle))
        self.assertEqual(json.loads(csv_row["occurrences"]), actual["events"][0]["occurrences"])
        self.assertEqual(csv_row["data_readiness"], actual["events"][0]["data_readiness"])
        self.assertEqual((exports / "latest/review.json").read_bytes(),
                         (exports / "gemini/latest/review.json").read_bytes())
        self.assertEqual(event_from_record(actual["events"][0]).occurrences, self.occurrences)
        self.assertEqual(self.repo.connection.execute("PRAGMA integrity_check").fetchone()[0], "ok")
        self.assertEqual(self.repo.connection.execute("PRAGMA foreign_key_check").fetchall(), [])

    def test_old_payload_migration_preserves_review_history_and_token(self):
        self.persist()
        token = self.repo.review_token("event", self.event.event_id, "one")
        self.assertEqual(self.repo.import_review(token, "event", self.event.event_id, "one",
                                                DECISIONS[1], "owner note"), "imported")
        row = self.repo.connection.execute("SELECT payload FROM event_snapshots").fetchone()[0]
        old = json.loads(row)
        for name in ("occurrences", "schedule_kind", "recurrence_rule"):
            old.pop(name)
        old_bytes = json.dumps(old)
        self.repo.connection.execute("UPDATE event_snapshots SET payload=?", (old_bytes,))
        self.repo.connection.execute("UPDATE events SET payload=?", (old_bytes,))
        document = self.repo.load()
        self.assertEqual(document["schedule_review"][0]["occurrences"], [])
        self.assertEqual(document["schedule_review"][0]["review_notes"], "owner note")
        self.assertEqual(self.repo.connection.execute("SELECT payload FROM event_snapshots").fetchone()[0],
                         old_bytes)
        self.persist("two", delta=1)
        self.assertEqual(self.repo.load()["events"][0]["review_notes"], "owner note")
        self.assertEqual(len(self.repo.list_runs()), 2)

    def test_workbook_sessions_readiness_and_review_roundtrip(self):
        self.persist()
        ExcelReviewExporter(self.book, self.repo).export(self.repo.load())
        workbook = load_workbook(self.book)
        sheet = workbook.active
        headers = [c.value for c in sheet[1]]
        value = lambda name: sheet.cell(2, headers.index(name) + 1)
        self.assertIn("مهر", value("جلسه‌ها").value)
        self.assertIn("تکمیل ظرفیت", value("جلسه‌ها").value)
        self.assertEqual(json.loads(value("occurrences").value)[0]["price_amount"], "150000")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertTrue(sheet.sheet_view.rightToLeft)
        value("تصمیم من").value = DECISIONS[1]
        value("یادداشت من").value = "keep me"
        workbook.save(self.book)
        workbook.close()
        self.assertEqual(ExcelReviewImporter(self.book, self.repo).import_reviews().imported, 1)
        self.persist("two", delta=1)
        ExcelReviewExporter(self.book, self.repo).export(self.repo.load())
        self.assertEqual(self.repo.load()["events"][0]["review_notes"], "keep me")
