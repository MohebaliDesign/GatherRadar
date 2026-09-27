from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from gatherradar.exports.excel import (ExcelReviewExporter, ExcelReviewImporter, TOKEN, literal, LOCAL_EVENT_VISIBLE, DUPLICATES)
from gatherradar.review.models import ReviewError
from gatherradar.review.schema import DECISIONS, DUPLICATE_DECISIONS
from gatherradar.review.workspace import local_workspace
from local_review_fakes import LocalCase


class ExcelTests(LocalCase):
    def setUp(self):
        super().setUp()
        self.persist()
        self.exporter = ExcelReviewExporter(self.book, self.repo)
        self.importer = ExcelReviewImporter(self.book, self.repo)
        self.exporter.export(self.repo.load())

    def edit(self, callback):
        workbook = load_workbook(self.book)
        try:
            callback(workbook)
            workbook.save(self.book)
        finally:
            workbook.close()

    def events(self, workbook):
        return next(s for s in workbook if s.title.startswith('رویدادها'))

    def cell(self, sheet, header, row=2):
        headers = [c.value for c in sheet[1]]
        return sheet.cell(row, headers.index(header) + 1)

    def test_rtl_columns_freeze_filter_dropdown_and_links(self):
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = self.events(workbook)
        self.assertEqual(tuple(c.value for c in sheet[1])[:len(LOCAL_EVENT_VISIBLE)], LOCAL_EVENT_VISIBLE)
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertTrue(sheet.auto_filter.ref)
        self.assertTrue(sheet.column_dimensions[get_column_letter(self.cell(sheet, 'event_id').column)].hidden)
        self.assertFalse(sheet.column_dimensions[get_column_letter(len(LOCAL_EVENT_VISIBLE))].hidden)
        self.assertIn(DECISIONS[2], sheet.data_validations.dataValidation[0].formula1)
        self.assertTrue(self.cell(sheet, 'منبع اصلی').hyperlink.target.startswith('https://'))
        self.assertEqual(workbook['_meta'].sheet_state, 'hidden')
        self.assertTrue(all(s.sheet_view.rightToLeft for s in workbook if s.title != '_meta'))
        self.assertTrue(any(s.title.startswith('مکان‌ها') for s in workbook))

    def test_jalali_and_weekday(self):
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = self.events(workbook)
        self.assertEqual(self.cell(sheet, 'تاریخ شمسی شروع').value, '۹ مهر ۱۴۰۵')
        self.assertEqual(self.cell(sheet, 'روز').value, 'پنجشنبه')

    def test_event_roundtrip_carry_forward_history_and_unrelated_identity(self):
        def change(workbook):
            sheet = self.events(workbook)
            self.cell(sheet, 'تصمیم من').value = DECISIONS[2]
            self.cell(sheet, 'یادداشت من').value = 'با دوستان'
            self.cell(sheet, 'عنوان').value = 'visual correction only'
        self.edit(change)
        self.assertEqual(self.importer.import_reviews().imported, 1)
        old = load_workbook(self.book)
        old_name = self.events(old).title
        old_values = list(old[old_name].values)
        old.close()
        self.persist('two', delta=60)
        self.exporter.export(self.repo.load())
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        self.assertEqual(list(workbook[old_name].values), old_values)
        new = next(s for s in workbook if s.title.startswith('رویدادها') and s.title != old_name)
        self.assertEqual(self.cell(new, 'تصمیم من').value, DECISIONS[2])
        self.assertEqual(self.cell(new, 'یادداشت من').value, 'با دوستان')
        self.assertEqual(self.cell(new, 'تصمیم من', 3).value, DECISIONS[0])
        self.assertNotEqual(self.cell(new, 'عنوان').value, 'visual correction only')

    def test_duplicate_roundtrip(self):
        def change(workbook):
            sheet = workbook[DUPLICATES]
            sheet['A2'] = DUPLICATE_DECISIONS[2]
            sheet['I2'] = 'دو برنامه جدا'
        self.edit(change)
        self.assertEqual(self.importer.import_reviews().imported, 1)
        self.persist('two', delta=2)
        self.exporter.export(self.repo.load())
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook[DUPLICATES]['I2'].value, 'دو برنامه جدا')
        self.assertEqual(workbook[DUPLICATES]['A2'].value, DUPLICATE_DECISIONS[2])

    def test_place_roundtrip_is_supported_only_for_exact_run(self):
        def change(workbook):
            sheet = next(s for s in workbook if s.title.startswith('مکان‌ها'))
            sheet['A2'] = DECISIONS[1]
            sheet['J2'] = 'مکان'
        self.edit(change)
        self.assertEqual(self.importer.import_reviews().imported, 1)
        self.assertEqual(self.repo.load()['places'][0]['review_notes'], 'مکان')
        self.persist('two', delta=2)
        self.assertEqual(self.repo.load()['places'][0]['review_notes'], '')

    def test_invalid_row_isolated_from_valid_row(self):
        def change(workbook):
            sheet = self.events(workbook)
            sheet['A2'] = 'invalid decision'
            sheet['A3'] = DECISIONS[1]
        self.edit(change)
        result = self.importer.import_reviews()
        self.assertEqual((result.imported, result.malformed), (1, 1))

    def test_changed_hidden_id_cannot_hijack_another_event(self):
        def change(workbook):
            sheet = self.events(workbook)
            self.cell(sheet, 'event_id').value = self.cell(sheet, 'event_id', 3).value
            sheet['A2'] = DECISIONS[2]
        self.edit(change)
        result = self.importer.import_reviews()
        self.assertEqual(result.imported, 0)
        self.assertGreater(result.malformed, 0)

    def test_unknown_id_rejected(self):
        self.edit(lambda w: setattr(self.cell(self.events(w), 'event_id'), 'value', 'unknown'))
        self.assertEqual(self.importer.import_reviews().malformed, 1)

    def test_formula_in_human_field_rejected(self):
        self.edit(lambda w: setattr(self.cell(self.events(w), 'یادداشت من'), 'value', '=1+1'))
        self.assertEqual(self.importer.import_reviews().malformed, 1)

    def test_literal_formula_looking_notes_import(self):
        self.edit(lambda w: literal(self.cell(self.events(w), 'یادداشت من'), '=literal note'))
        self.assertEqual(self.importer.import_reviews().imported, 1)
        self.assertIn('=literal note', [e['review_notes'] for e in self.repo.load()['events']])

    def test_formula_prefixes_stay_literal_in_generated_workbook(self):
        for prefix in ('=', '+', '-', '@'):
            event = replace(self.result.events[0], title=prefix + 'unsafe')
            run_id = 'formula_' + str(ord(prefix))
            self.persist(run_id, replace(self.result, events=(event,)), delta=ord(prefix))
            path = self.root / (run_id + '.xlsx')
            ExcelReviewExporter(path, self.repo).export(self.repo.load(run_id))
            workbook = load_workbook(path, data_only=False)
            try:
                cell = self.cell(self.events(workbook), 'عنوان')
                self.assertEqual(cell.value, prefix + 'unsafe')
                self.assertEqual(cell.data_type, 's')
            finally:
                workbook.close()

    def test_corruption_and_newer_schema_never_overwritten(self):
        for content in (b'not an xlsx', b'PKbroken'):
            self.book.write_bytes(content)
            with self.assertRaises(ReviewError):
                self.exporter.export(self.repo.load())
            self.assertEqual(self.book.read_bytes(), content)

    def test_newer_schema_import_refused(self):
        self.edit(lambda w: setattr(w['_meta']['B1'], 'value', 99))
        with self.assertRaises(ReviewError):
            self.importer.import_reviews()

    def test_missing_review_header_refused(self):
        self.edit(lambda w: setattr(self.events(w)['A1'], 'value', 'missing'))
        with self.assertRaises(ReviewError):
            self.importer.import_reviews()

    def test_atomic_replace_failure_preserves_old_bytes_and_database(self):
        before = self.book.read_bytes()
        with patch.object(Path, 'replace', side_effect=PermissionError('locked')):
            with self.assertRaisesRegex(ReviewError, 'Close it'):
                self.exporter.export(self.repo.load())
        self.assertEqual(self.book.read_bytes(), before)
        self.assertTrue(self.repo.find_run('one'))
        self.assertEqual(list(self.book.parent.glob('*.tmp.xlsx')), [])

    def test_same_run_reexport_preserves_edits_without_duplicate_snapshot(self):
        self.edit(lambda w: setattr(self.events(w)['B2'], 'value', 'local edit'))
        self.exporter.export(self.repo.load())
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        self.assertEqual(sum(s.title.startswith('رویدادها') for s in workbook), 1)
        self.assertEqual(self.events(workbook)['B2'].value, 'local edit')

    def test_same_minute_names_fit_xlsx_limit_and_are_unique(self):
        self.persist('two', delta=1)
        self.exporter.export(self.repo.load())
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        self.assertEqual(sum(s.title.startswith('رویدادها') for s in workbook), 2)
        self.assertTrue(all(len(name) <= 31 for name in workbook.sheetnames))

    def test_row_and_tab_order_cannot_roll_back_human_state(self):
        self.edit(lambda w: setattr(self.events(w)['A2'], 'value', DECISIONS[1]))
        self.importer.import_reviews()
        self.persist('two', delta=60)
        self.exporter.export(self.repo.load())
        self.assertEqual(self.importer.import_reviews().imported, 0)
        self.assertEqual(sum(e['review_decision'] == DECISIONS[1] for e in self.repo.load()['events']), 1)

    def test_invalid_rows_hold_workbook_but_portable_outputs_continue(self):
        self.edit(lambda w: setattr(self.events(w)['A2'], 'value', 'bad'))
        before = self.book.read_bytes()
        workspace = local_workspace(self.repo, self.book, self.root / 'exports')
        workspace.import_reviews()
        paths, failures = workspace.export('one')
        self.assertTrue(failures)
        self.assertEqual(self.book.read_bytes(), before)
        self.assertTrue(any(path.name == 'review.json' for path in paths))

    def test_overlong_text_fails_without_silent_truncation(self):
        from openpyxl import Workbook
        workbook = Workbook()
        with self.assertRaises(ReviewError):
            literal(workbook.active['A1'], 'x' * 32768)

    def test_all_supporting_source_urls_have_clickable_cells(self):
        links = ('https://example.test/a', 'https://example.test/b')
        event = replace(self.result.events[0], evidence_urls=links, canonical_source_url=links[0])
        self.persist('links', replace(self.result, events=(event,)), delta=1)
        path = self.root / 'links.xlsx'
        ExcelReviewExporter(path, self.repo).export(self.repo.load('links'))
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = self.events(workbook)
        self.assertEqual(self.cell(sheet, 'پیوند منبع ۱').hyperlink.target, links[0])
        self.assertEqual(self.cell(sheet, 'پیوند منبع ۲').hyperlink.target, links[1])
        self.assertEqual(ExcelReviewImporter(path, self.repo).import_reviews().malformed, 0)

    def test_unknown_precision_sorts_after_dates_and_time_sort_is_preserved(self):
        from gatherradar.domain.temporal import DatePrecision
        unknown = replace(self.result.events[0], date_precision=DatePrecision.UNKNOWN, start_date=None,
                          starts_at=None, start_time=None)
        self.persist('sorted', replace(self.result, events=(unknown, self.result.events[1])), delta=2)
        path = self.root / 'sorted.xlsx'
        ExcelReviewExporter(path, self.repo).export(self.repo.load('sorted'))
        workbook = load_workbook(path)
        self.addCleanup(workbook.close)
        sheet = self.events(workbook)
        self.assertEqual(self.cell(sheet, 'event_id', 3).value, unknown.event_id)
