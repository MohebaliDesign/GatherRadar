"""Source facts and session presentation survive the real local export boundary."""
import csv
from datetime import datetime
from openpyxl import load_workbook
from gatherradar.deduplication.models import CanonicalizationResult
from gatherradar.review.serialization import review_window
from gatherradar.review.workspace import local_workspace
from local_review_fakes import LocalCase
from test_event_recall import pipeline, REFERENCE


class FieldPersistenceTests(LocalCase):
    def export(self, text, reference=REFERENCE):
        events, _ = pipeline(text)
        self.assertEqual(len(events), 1)
        self.repo.sync(CanonicalizationResult(events=events), review_window(events, reference),
            run_id='fields', started_at=reference, finished_at=reference, days=14,
            review_timezone='Asia/Tehran', source_count=1)
        self.workspace = local_workspace(self.repo, self.book, self.root/'exports')
        self.assertFalse(self.workspace.export('fields')[1])
        return events[0], self.repo.load('fields')

    def test_range_location_and_hours_reach_excel_without_invented_address(self):
        event, document = self.export('نمایشگاه انفرادی «آبی» از هنرمند نمونه\n'
            '۲۴ مهر تا ۱۵ آبان ۱۴۰۵\nساعت ۱۳ تا ۲۱ (شنبه‌ها تعطیل)\n'
            '📍 محله نمونه، مجتمع سپید، گالری ابر/باران')
        row = document['events'][0]
        self.assertEqual(row['location_context'], 'مجتمع سپید')
        self.assertEqual(row['venue_name'], 'گالری ابر/باران')
        self.assertEqual(row['area_text'], 'محله نمونه')
        self.assertIsNone(row['address'])
        self.assertIsNone(row['city'])
        self.assertIsNone(row['next_session_date'])
        self.assertEqual(row['occurrences'][0]['end_date'], '2026-11-06')
        self.assertIn('شنبه', row['source_schedule_text'])
        self.assertIn('location_context', {p.field for p in event.field_provenance})
        book = load_workbook(self.book)
        visible = dict(zip(next(book.active.values), list(book.active.values)[1]))
        self.assertIn('مجتمع سپید', visible.values())
        self.assertIn('گالری ابر/باران', visible.values())
        self.assertIn('محله نمونه', visible.values())
        book.close()
        for folder in ('latest', 'gemini/latest'):
            with (self.root/'exports'/folder/'events.csv').open(encoding='utf-8-sig', newline='') as stream:
                self.assertEqual(next(csv.DictReader(stream))['location_context'], 'مجتمع سپید')

    def test_next_session_uses_stored_clock_and_keeps_parent_and_history(self):
        event, document = self.export('کارگاه سفال آبی\n۳۰ مهر ساعت ۱۶ الی ۲۰، ۱ آبان ساعت ۱۲ الی ۲۰\nثبت نام')
        row = document['events'][0]
        self.assertEqual(row['next_session_date'], '2026-10-23')
        self.assertEqual(row['next_session_start_time'], '12:00:00')
        self.assertEqual(row['next_session_end_time'], '20:00:00')
        self.assertEqual(row['remaining_session_count'], 1)
        self.assertEqual(row['start_date'], event.start_date.isoformat() if event.start_date else None)
        self.assertEqual(len(event.occurrences), 2)
        book = load_workbook(self.book)
        values = dict(zip(next(book.active.values), list(book.active.values)[1]))
        self.assertEqual(values['شروع جلسه بعدی'], '۱۲:۰۰')
        self.assertEqual(values['جلسه‌های باقی‌مانده'], 1)
        names = book.sheetnames
        book.close()
        self.assertFalse(self.workspace.export('fields')[1])
        self.assertEqual(self.repo.load('fields'), document)
        book = load_workbook(self.book)
        self.assertEqual(book.sheetnames, names)
        book.close()

    def test_session_already_started_is_not_called_next(self):
        _, document = self.export('کارگاه سفال آبی\n۱ آبان ساعت ۱۲ الی ۲۰\nثبت نام',
            datetime.fromisoformat('2026-10-23T14:00:00+03:30'))
        self.assertIsNone(document['events'][0]['next_session_date'])
        self.assertEqual(document['events'][0]['remaining_session_count'], 0)

    def test_description_only_reprojection_preserves_previous_sheet(self):
        from copy import deepcopy
        from gatherradar.exports.excel import ExcelReviewExporter
        _, document = self.export('کارگاه سفال آبی\n۱ آبان ساعت ۱۲ الی ۲۰\nثبت نام')
        book = load_workbook(self.book)
        old_name, old_rows = book.active.title, list(book.active.values)
        book.close()
        revised = deepcopy(document)
        revised['events'][0]['description_text'] += '\nتوضیح تکمیلی منبع'
        exporter = ExcelReviewExporter(self.book, self.repo)
        exporter.export(revised)
        book = load_workbook(self.book)
        self.assertNotEqual(book.active.title, old_name)
        self.assertEqual(list(book[old_name].values), old_rows)
        self.assertIn('توضیح تکمیلی منبع', str(list(book.active.values)))
        names = book.sheetnames
        book.close()
        exporter.export(revised)
        book = load_workbook(self.book)
        self.assertEqual(book.sheetnames, names)
        book.close()
