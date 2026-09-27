import builtins
import io
import json
from contextlib import redirect_stdout, redirect_stderr
from dataclasses import replace
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from gatherradar.cli import build_parser, main
from gatherradar.config import load_sources
from gatherradar.orchestration.refresh import RefreshService
from gatherradar.orchestration.website_run import website_output_path
from gatherradar.review.schema import DECISIONS
from gatherradar.review.workspace import local_workspace
from gatherradar.storage import JsonlRawItemStore
from local_review_fakes import LocalCase, STAMP
from test_refresh import raw, CONFIG


class LocalRefreshTests(LocalCase):
    def setUp(self):
        super().setUp()
        self.source = next(s for s in load_sources(CONFIG) if s.id == 'davvvat_website')
        self.item = raw(self.source, captured=STAMP)
        self.workspace = local_workspace(self.repo, self.book, self.root / 'exports')
        self.collector = Mock(return_value=SimpleNamespace(observed_items=(self.item,),
            observed_item_ids=(self.item.id,), failed=0, new=0, changed=0, already_existing=1))
        self.service = RefreshService(self.repo, website_collect=self.collector,
            instagram_collect=Mock(side_effect=RuntimeError('secret')), now=lambda: STAMP,
            import_reviews=self.workspace.import_reviews, export_reviews=self.workspace.export)

    def refresh(self, **kwargs):
        return self.service.run(kwargs.pop('sources', (self.source.id,)), config_path=CONFIG,
            data_dir=self.root, run_id=kwargs.pop('run_id', 'local'), **kwargs)

    def test_default_refresh_without_google_installed(self):
        original = builtins.__import__
        def blocked(name, *args, **kwargs):
            if name.startswith(('google', 'httplib2')):
                raise ModuleNotFoundError('blocked optional dependency')
            return original(name, *args, **kwargs)
        with patch('builtins.__import__', side_effect=blocked):
            self.assertEqual(build_parser().parse_args(['refresh', '--all-enabled']).days, 14)
            result = self.refresh()
            self.assertEqual(result.export_failures, ())
            self.assertEqual(result.snapshot.events, 1)
            self.assertTrue(self.book.exists())
            self.assertTrue((self.root / 'exports/gemini/latest/manifest.json').exists())

    def test_partial_sources_preserve_success_and_exact_membership(self):
        stale = raw(self.source, 'stale')
        JsonlRawItemStore(website_output_path(self.source, self.root)).append_new((stale,))
        result = self.refresh(sources=(self.source.id,'davvvat_instagram'))
        self.assertEqual(result.snapshot.status, 'partial')
        document = self.repo.load()
        self.assertEqual(document['events'][0]['source_item_ids'], [self.item.id])
        self.assertEqual(document['run']['sources'][1]['observed_item_ids'], [self.item.id])
        self.assertNotIn('secret', json.dumps(document))

    def test_all_sources_failed_are_persisted(self):
        result = self.refresh(sources=('davvvat_instagram',))
        self.assertEqual(result.snapshot.status, 'failed')
        self.assertEqual(self.repo.load()['run']['status'], 'failed')
        self.assertEqual(self.repo.load()['events'], [])

    def test_export_failure_does_not_rollback_or_mark_run_failed(self):
        self.service.export_reviews = Mock(side_effect=RuntimeError('private path'))
        result = self.refresh()
        self.assertTrue(result.export_failures)
        self.assertEqual(self.repo.load()['run']['status'], 'success')
        self.assertNotIn('private path', str(result))
        self.service.export_reviews = self.workspace.export
        second = self.refresh()
        self.assertEqual(second.export_failures, ())
        self.assertEqual(self.collector.call_count, 1)

    def test_automatic_import_precedes_collection_and_carries_state(self):
        self.refresh()
        workbook = load_workbook(self.book)
        sheet = next(s for s in workbook if s.title.startswith('رویدادها'))
        sheet['A2'] = DECISIONS[2]
        headers = [c.value for c in sheet[1]]
        sheet.cell(2,headers.index('یادداشت من')+1).value = 'automatic import'
        workbook.save(self.book)
        workbook.close()
        original = self.collector.return_value
        def collect(*args, **kwargs):
            self.assertEqual(self.repo.load()['events'][0]['review_notes'], 'automatic import')
            return original
        self.collector.side_effect = collect
        result = self.refresh(run_id='next')
        self.assertEqual(result.import_summary.imported, 1)
        self.assertEqual(self.repo.load('next')['events'][0]['review_decision'], DECISIONS[2])

    def test_offline_reexport_uses_sqlite_only(self):
        self.refresh()
        other = local_workspace(self.repo, self.root / 'rebuilt/book.xlsx', self.root / 'rebuilt/exports')
        with patch('gatherradar.orchestration.refresh.run_canonical_review', side_effect=AssertionError('analysis')):
            paths, failures = other.export('latest')
        self.assertEqual(failures, ())
        self.assertEqual(len(paths), 21)
        self.assertEqual(self.collector.call_count, 1)

    def test_explicit_import_and_status_commands_are_offline(self):
        self.refresh()
        args = ['--db', str(self.repo.path), '--workbook', str(self.book)]
        with redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            self.assertEqual(main(['review', 'import', *args]), 0)
            self.assertEqual(main(['status', *args]), 0)

    def test_google_command_missing_dependencies_has_actionable_message(self):
        original = builtins.__import__
        def blocked(name, *args, **kwargs):
            if name.startswith(('google','httplib2')):
                raise ModuleNotFoundError('no google')
            return original(name, *args, **kwargs)
        output = io.StringIO()
        with patch('builtins.__import__', side_effect=blocked), redirect_stdout(io.StringIO()), redirect_stderr(output):
            self.assertEqual(main(['sheets', 'status', '--data-dir', str(self.root)]), 1)
        self.assertIn('.[google-sheets]', output.getvalue())
        self.assertNotIn('Traceback', output.getvalue())

    def test_csv_json_gemini_cli_without_xlsx_or_google_packages(self):
        self.refresh()
        original = builtins.__import__
        def blocked(name, *args, **kwargs):
            if name.startswith(('google','httplib2','openpyxl')):
                raise ModuleNotFoundError(name)
            return original(name, *args, **kwargs)
        with patch('builtins.__import__', side_effect=blocked), redirect_stdout(io.StringIO()):
            for kind in ('csv','json','gemini'):
                self.assertEqual(main(['export', kind, '--db', str(self.repo.path), '--export-dir', str(self.root / 'other')]), 0)
