"""Embedded sessions across discovery, canonical storage and supported exports."""
import csv
import json
import sqlite3
from contextlib import closing
from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook

from gatherradar.collectors.website import WebsiteCollector
from gatherradar.collectors.websites.transport import Page
from gatherradar.deduplication import CandidateContext, canonicalize
from gatherradar.domain import Source, SourceType
from gatherradar.domain.website import WebsiteConfig
from gatherradar.exports.excel import ExcelReviewExporter
from gatherradar.exports.gemini import GeminiBundleExporter
from gatherradar.exports.portable import CsvReviewExporter, JsonReviewExporter
from gatherradar.extraction import DiscoveryService, RuleBasedDiscoveryProvider
from gatherradar.extraction.service import build_extraction_input
from gatherradar.normalization import NormalizationService
from gatherradar.review.readiness import assess_readiness, DataReadiness
from gatherradar.review.records import event_from_record
from local_review_fakes import LocalCase
from test_jabama_sessions import BASE, URL, SESSION, _page

SOURCE = Source('jabama_test', 'jabama', 'Jabama', SourceType.WEBSITE, BASE,
                timezone='Asia/Tehran', website=WebsiteConfig('jabama-events', ('/events/',)))
STAMP = datetime(2026, 9, 26, tzinfo=timezone.utc)


def collected(page):
    class Transport:
        def fetch(self, url, *, allowed=None):
            return Page(url, page if url == URL else '<a href="/events/demo-event">برنامه</a>')
    raw, = WebsiteCollector(transport=Transport(), now=lambda: STAMP).collect(SOURCE, limit=1).items
    return raw


def pipeline(raw):
    discovery = DiscoveryService(RuleBasedDiscoveryProvider()).discover(raw, SOURCE)
    if discovery.event is None:
        raise AssertionError(f'Expected supported session Event: {discovery}')
    outcome = NormalizationService().normalize(discovery.event, raw, SOURCE)
    result = canonicalize((CandidateContext(outcome, discovery.event, raw, SOURCE, STAMP),))
    if not result.events:
        raise AssertionError(f'Expected canonical Event: {result}')
    return discovery.event, result


class JabamaSessionPipelineTests(LocalCase):
    def test_visible_two_day_schedule_and_one_embedded_session_do_not_duplicate_overlap(self):
        # Verified saved public page shape: visible two-day schedule, while the
        # embedded list retains only the second date's bookable session.
        raw = collected(_page().replace('<h1>کارگاه نمونه</h1>',
                                        '<h1>کارگاه نمونه</h1><p>۱ و ۲ مهر ماه ساعت ۲۰:۳۰</p>'))
        _, result = pipeline(raw)
        event, = result.events
        self.assertEqual(len(event.occurrences), 2)
        values = sorted(event.occurrences, key=lambda o: o.start_date)
        self.assertEqual([o.start_date.isoformat() for o in values], ['2026-09-23', '2026-09-24'])
        self.assertIsNone(values[0].source_session_id)
        self.assertIsNone(values[0].remaining_seats)
        self.assertEqual(values[1].source_session_id, 'demo-session-a')
        self.assertEqual(values[1].remaining_seats, 45)
        self.assertEqual(values[1].end_time.isoformat(), '22:30:00')
        self.assertIsNone(event.start_date)

    def test_session_only_place_like_page_is_an_event_with_concrete_occurrence(self):
        page = _page().replace('کارگاه نمونه', 'کافه گالری آبی').replace(
            'ساخت یک اثر با راهنمایی مربی.', 'کافه گالری آبی فضای آرام و زیبایی برای نوشیدن قهوه دارد.')
        raw = collected(page)
        candidate, result = pipeline(raw)
        self.assertEqual(len(candidate.source_sessions), 1)
        event, = result.events
        value, = event.occurrences
        self.assertEqual(value.start_time.isoformat(), '20:30:00')
        self.assertEqual(value.end_time.isoformat(), '22:30:00')
        self.assertEqual(value.remaining_seats, 45)
        self.assertEqual(value.evidence_urls, (URL,))
        self.assertEqual(value.source_item_ids, (raw.id,))
        self.assertIn('/sessions/0', value.source_locator)
        self.assertIsNone(value.price_amount)  # page starting price is not session evidence
        self.assertEqual(assess_readiness(event).status, DataReadiness.READY)

    def test_multiple_sessions_survive_sqlite_xlsx_csv_json_and_gemini(self):
        first = {**SESSION, 'startIso': '2026-09-27T17:00:00Z', 'pricePerPerson': 250000}
        second = {**SESSION, 'id': 'demo-session-b', 'startIso': '2026-09-28T17:00:00Z',
                  'pricePerPerson': 300000, 'hostDiscountPct': 10, 'remainingSeats': 0}
        raw = collected(_page([first, second]))
        _, result = pipeline(raw)
        event, = result.events
        self.assertEqual(len(event.occurrences), 2)
        self.assertIsNone(event.start_date)
        a, b = sorted(event.occurrences, key=lambda o: o.start_date)
        self.assertEqual(a.price_amount, Decimal(250000))
        self.assertIsNone(b.price_amount)
        self.assertIsNone(b.price_text)
        self.assertEqual(b.price_per_person, '300000')
        self.assertEqual(b.host_discount_pct, '10')
        self.assertEqual(b.remaining_seats, 0)
        self.assertIn('normalization:embedded_session_discount_unresolved',
                      {d.code for d in event.diagnostics})
        self.persist(result=result)
        document = self.repo.load()
        stored, = document['events']
        restored = event_from_record(stored)
        self.assertEqual(restored.occurrences, event.occurrences)
        paths = self.root / 'exports'
        ExcelReviewExporter(self.book, self.repo).export(document)
        CsvReviewExporter(paths).export(document)
        JsonReviewExporter(paths).export(document)
        GeminiBundleExporter(paths / 'gemini').export(document)
        json_record = json.loads((paths / 'latest/review.json').read_text(encoding='utf-8'))['events'][0]
        with (paths / 'latest/events.csv').open(encoding='utf-8-sig', newline='') as handle:
            csv_record = next(csv.DictReader(handle))
        gemini_record = json.loads((paths / 'gemini/latest/review.json').read_text(encoding='utf-8'))['events'][0]
        self.assertEqual(json_record['occurrences'], stored['occurrences'])
        self.assertEqual(json.loads(csv_record['occurrences']), stored['occurrences'])
        self.assertEqual(gemini_record['occurrences'], stored['occurrences'])
        with closing(sqlite3.connect(self.repo.path)) as connection:
            self.assertEqual(connection.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
            self.assertEqual(connection.execute('PRAGMA foreign_key_check').fetchall(), [])
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = next(s for s in workbook if s.title.startswith('رویدادها'))
        headers = [c.value for c in sheet[1]]
        schedule = sheet.cell(2, headers.index('جلسه‌ها') + 1).value
        self.assertIn('۲۰:۳۰', schedule)
        self.assertNotIn('hostDiscountPct', schedule)
        self.assertNotIn('pricePerPerson', schedule)

    def test_untrusted_session_metadata_cannot_crash_or_borrow_other_event_data(self):
        raw = collected(_page())
        metadata = {**raw.raw_metadata, 'source_session_diagnostics': [{}, [], None, 'embedded_sessions_empty']}
        value = build_extraction_input(replace(raw, raw_metadata=metadata), SOURCE)
        self.assertEqual(value.source_session_diagnostics, ('embedded_sessions_empty',))
        poisoned = {**metadata, 'source_sessions': [
            {**raw.raw_metadata['source_sessions'][0], 'source_event_id': 'another-event'}]}
        value = build_extraction_input(replace(raw, raw_metadata=poisoned), SOURCE)
        self.assertEqual(value.source_sessions, ())
        self.assertIn('embedded_session_identity_conflict', value.source_session_diagnostics)

    def test_partially_unrecognized_session_schema_cannot_be_ready(self):
        raw = collected(_page([SESSION, {**SESSION, 'id': 'unsupported', 'timeslots': ['10:00']}]))
        _, result = pipeline(raw)
        event, = result.events
        self.assertEqual(len(event.occurrences), 1)
        self.assertIn('normalization:embedded_session_schema_drift', {d.code for d in event.diagnostics})
        self.assertNotEqual(assess_readiness(event).status, DataReadiness.READY)
