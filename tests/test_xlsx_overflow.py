"""Lossless XLSX export of values above the 32,767-character cell limit.

Live run ca3548cb65534d868277f140cdc92855 failed because one Jabama Event with 50
embedded sessions serialized ~56,591 characters of occurrences into one cell.
Synthetic sessions below mirror that record's shape; no runtime data is used.
"""
from __future__ import annotations

import csv
import json
import math
from dataclasses import replace
from datetime import date, datetime, time, timedelta, timezone

from openpyxl import Workbook, load_workbook

from gatherradar.domain.occurrence import EventOccurrence, ScheduleKind, identify_occurrence
from gatherradar.domain.temporal import DatePrecision
from gatherradar.exports.excel import (DESCRIPTION, OVERFLOW, OVERFLOW_CHUNK, OVERFLOW_HEADERS, TRUNCATED,
                                       XLSX_CELL_LIMIT, CellLimitError, ExcelReviewExporter,
                                       ExcelReviewImporter, fit_rows, open_workspace, overflow_value, table)
from gatherradar.review.models import ReviewError
from gatherradar.review.records import dumps, primitive
from gatherradar.review.schema import DECISIONS
from gatherradar.review.workspace import local_workspace
from local_review_fakes import LocalCase
from test_source_fields_outputs import result


def sessions(count: int, *, per_day: int = 1, unverified: bool = True) -> tuple[EventOccurrence, ...]:
    """Unverified-shape sessions (the live case) go to Schedule Review with every
    session; verified ones within the 14-day window are actionable."""
    values = []
    for index in range(count):
        day = date(2026, 9, 27) + timedelta(days=index // per_day)
        start = datetime.combine(day, time(8 + index % per_day), tzinfo=timezone.utc)
        values.append(identify_occurrence(EventOccurrence(
            'pending', day, 'UTC', start_time=start.time(), end_time=time(20), starts_at=start,
            ends_at=start.replace(hour=20), date_precision=DatePrecision.EXACT,
            source_schedule_text=start.isoformat(), source_session_id=f'{index:08d}',
            source_locator=f'react-router:loaderData/routes/events.$id/sessions/{index}',
            source_session_origin='example.test', source_item_ids=('website:example:legacy-941',),
            candidate_ids=('candidate:' + 'c' * 64,), evidence_urls=('https://example.test/events/legacy-941',),
            evidence_fragment_ids=('evidence:' + 'a' * 64, 'evidence:' + 'b' * 64),
            diagnostics=(('embedded_session_unverified_shape',) if unverified
                         else ('embedded_session_corroborated_by_page_text',))), 'event:synthetic'))
    return tuple(values)


class OverflowExportCase(LocalCase):
    def persist_sessions(self, count: int, run_id: str = 'big', delta: int = 0, *, per_day: int = 1,
                         unverified: bool = True, **fields):
        base = result(**fields)
        event = replace(base.events[0], occurrences=sessions(count, per_day=per_day, unverified=unverified), schedule_kind=ScheduleKind.DISCRETE,
                        start_date=None, start_time=None, starts_at=None, date_precision=DatePrecision.UNKNOWN)
        self.result = replace(base, events=(event,))
        self.persist(run_id, delta=delta)
        return self.repo.load(run_id)

    def export(self, run_id: str):
        workspace = local_workspace(self.repo, self.book, self.root / 'exports')
        paths, warnings = workspace.export(run_id)
        self.assertEqual(warnings, ())
        return paths

    def event_sheet(self, workbook):
        """The sheet holding the single Event: Schedule Review for unverified sessions."""
        return next(s for s in workbook if s.title.startswith(('رویدادها', 'زمان‌بندی')) and s.max_row > 1)

    def record(self, document):
        return (document['events'] or document['schedule_review'])[0]


class LargeOccurrenceExportTests(OverflowExportCase):
    def check_lossless(self, count: int, **options):
        document = self.persist_sessions(count, **options)
        event = self.record(document)
        expected = dumps(primitive(event['occurrences']))
        self.assertGreater(len(expected), XLSX_CELL_LIMIT)
        self.export('big')
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = self.event_sheet(workbook)
        headers = [c.value for c in sheet[1]]
        cell = sheet.cell(2, headers.index('occurrences') + 1).value
        self.assertTrue(cell.startswith('[[gatherradar-overflow sha256='))
        self.assertIn(f'length={len(expected)}', cell)
        restored = overflow_value(workbook, sheet.title, event['event_id'], 'occurrences')
        self.assertEqual(restored, expected)  # byte-for-byte, nothing dropped
        self.assertEqual(len(json.loads(restored)), count)
        self.assertEqual(workbook[OVERFLOW].sheet_state, 'hidden')
        self.assertEqual(workbook[OVERFLOW].max_row - 1, math.ceil(len(expected) / OVERFLOW_CHUNK))
        self.assertTrue(all(len(v) <= XLSX_CELL_LIMIT for s in workbook for row in s.iter_rows(values_only=True)
                            for v in row if isinstance(v, str)))
        self.assertEqual(sheet.max_row - 1, 1)  # user-facing Event count unchanged
        self.assertEqual(len(document['events']) + len(document['schedule_review']), 1)
        return document, workbook

    def test_fifty_occurrences_like_the_live_jabama_event(self):
        self.check_lossless(50)

    def test_hundred_occurrences(self):
        self.check_lossless(100)

    def test_many_actionable_sessions_inside_the_window(self):
        document, workbook = self.check_lossless(100, per_day=8, unverified=False)
        self.assertEqual(len(document['events']), 1)
        self.assertTrue(self.event_sheet(workbook).title.startswith('رویدادها'))

    def test_sqlite_csv_json_and_workbook_agree_on_every_occurrence(self):
        document, workbook = self.check_lossless(50)
        event_id = self.record(document)['event_id']
        latest = self.root / 'exports' / 'latest'
        with (latest / 'schedule_review.csv').open(encoding='utf-8-sig', newline='') as stream:
            row = next(r for r in csv.DictReader(stream) if r['event_id'] == event_id)
        exported = json.loads((latest / 'review.json').read_text(encoding='utf-8'))
        counts = {len(self.record(document)['occurrences']), len(json.loads(row['occurrences'])),
                  len(self.record(exported)['occurrences']),
                  len(json.loads(overflow_value(workbook, self.event_sheet(workbook).title, event_id, 'occurrences')))}
        self.assertEqual(counts, {50})

    def test_very_long_visible_session_list_gets_an_explicit_marker_and_full_copy(self):
        document = self.persist_sessions(1500)
        self.export('big')
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = self.event_sheet(workbook)
        headers = [c.value for c in sheet[1]]
        visible = sheet.cell(2, headers.index('جلسه‌ها') + 1).value
        self.assertLessEqual(len(visible), XLSX_CELL_LIMIT)
        self.assertTrue(visible.endswith('[متن کامل در برگه پنهان _overflow و در review.json]'))
        full = overflow_value(workbook, sheet.title, self.record(document)['event_id'], 'جلسه‌ها')
        self.assertEqual(len(full.splitlines()), 1500)
        self.assertTrue(full.startswith(visible[:100]))

    def test_one_oversized_technical_field_and_description(self):
        text = 'متن منبع ' * 5000
        document = self.persist_sessions(2, description_text=text, source_schedule_text='۸ آبان ' * 6000)
        self.export('big')
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = self.event_sheet(workbook)
        headers = [c.value for c in sheet[1]]
        event_id = self.record(document)['event_id']
        self.assertTrue(sheet.cell(2, headers.index(DESCRIPTION) + 1).value.endswith(TRUNCATED))
        self.assertEqual(overflow_value(workbook, sheet.title, event_id, DESCRIPTION), text)
        self.assertEqual(overflow_value(workbook, sheet.title, event_id, 'source_schedule_text'),
                         '۸ آبان ' * 6000)


class OverflowWorkspaceTests(OverflowExportCase):
    def test_import_round_trip_and_same_run_reexport_are_stable(self):
        document = self.persist_sessions(50)
        self.export('big')
        workbook = load_workbook(self.book)
        sheet = self.event_sheet(workbook)
        headers = [c.value for c in sheet[1]]
        sheet.cell(2, headers.index('تصمیم من') + 1).value = DECISIONS[2]
        sheet.cell(2, headers.index('یادداشت من') + 1).value = 'یادداشت بلندتر از معمول'
        workbook.save(self.book)
        workbook.close()
        summary = ExcelReviewImporter(self.book, self.repo).import_reviews()
        self.assertEqual((summary.imported, summary.malformed, summary.conflicts), (1, 0, 0))
        stored = self.record(self.repo.load('big'))
        self.assertEqual((stored['review_decision'], stored['review_notes']),
                         (DECISIONS[2], 'یادداشت بلندتر از معمول'))
        self.assertEqual(stored['occurrences'], self.record(document)['occurrences'])  # facts untouched
        before = load_workbook(self.book)
        names, overflow_rows = before.sheetnames, list(before[OVERFLOW].values)
        before.close()
        self.export('big')  # same run, same projection: no new sheet, no new overflow rows
        after = load_workbook(self.book)
        self.addCleanup(after.close)
        self.assertEqual(after.sheetnames, names)
        self.assertEqual(list(after[OVERFLOW].values), overflow_rows)
        self.assertEqual(ExcelReviewImporter(self.book, self.repo).import_reviews().malformed, 0)

    def test_earlier_workbooks_stay_readable_and_history_is_not_rewritten(self):
        self.persist('small')
        self.export('small')  # pre-change shape: no _overflow sheet
        old = load_workbook(self.book)
        self.assertNotIn(OVERFLOW, old.sheetnames)
        old_sheet = next(s for s in old if s.title.startswith('رویدادها'))
        old_title, old_values = old_sheet.title, list(old_sheet.values)
        old.close()
        self.persist_sessions(50, run_id='big', delta=10)
        self.export('big')
        workbook = open_workspace(self.book)
        self.addCleanup(workbook.close)
        self.assertEqual(list(workbook[old_title].values), old_values)
        self.assertIn(OVERFLOW, workbook.sheetnames)
        self.assertEqual(workbook.sheetnames[-1], OVERFLOW)


class OverflowGuardTests(OverflowExportCase):
    def test_reviewer_owned_values_are_never_truncated(self):
        headers = ('تصمیم من', 'یادداشت من', 'event_id')
        with self.assertRaises(CellLimitError) as caught:
            fit_rows('sheet', headers, [[DECISIONS[0], 'ن' * 40000, 'event:abc']], 'event_id', 2)
        message = str(caught.exception)
        for part in ('event:abc', 'یادداشت من', '40000', 'sheet=sheet'):
            self.assertIn(part, message)

    def test_unfitted_cells_report_sheet_record_field_and_length(self):
        sheet = Workbook().active
        sheet.title = 'queue'
        with self.assertRaises(CellLimitError) as caught:
            table(sheet, ('pair_key', 'reason'), [['pair:1', 'x' * 33000]], 2)
        self.assertIn('sheet=queue, record_id=pair:1, field=reason', str(caught.exception))
        self.assertIn('33000', str(caught.exception))

    def test_altered_or_missing_chunks_are_refused(self):
        document = self.persist_sessions(50)
        self.export('big')
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        title = self.event_sheet(workbook).title
        event_id = self.record(document)['event_id']
        workbook[OVERFLOW].cell(2, OVERFLOW_HEADERS.index('text') + 1).value = 'altered'
        with self.assertRaises(ReviewError):
            overflow_value(workbook, title, event_id, 'occurrences')
        with self.assertRaises(ReviewError):
            overflow_value(workbook, title, event_id, 'diagnostics')
