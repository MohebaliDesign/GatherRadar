"""New source-supported fields across SQLite, XLSX, CSV, JSON, Gemini and Sheets rows."""
import csv
import json
import sqlite3
from dataclasses import replace
from hashlib import sha256

from openpyxl import load_workbook

from gatherradar.deduplication import canonicalize
from gatherradar.domain import EventStatus
from gatherradar.domain.temporal import DatePrecision
from gatherradar.exports.excel import (DESCRIPTION, TRUNCATED, WORKBOOK_VERSION, XLSX_CELL_LIMIT,
                                       ExcelReviewExporter, ExcelReviewImporter)
from gatherradar.exports.excel_design import ROW_MAX
from gatherradar.exports.gemini import GeminiBundleExporter, INSTRUCTIONS
from gatherradar.exports.portable import CsvReviewExporter, JsonReviewExporter
from gatherradar.review.models import ReviewError
from gatherradar.review.records import SCHEMA_VERSION, event_from_record
from gatherradar.review.schema import DECISIONS, EVENT_HEADERS, EVENT_VISIBLE
from gatherradar.review.serialization import event_row
from gatherradar.storage.sqlite_schema import SCHEMA_VERSION as DB_VERSION
from deduplication_fakes import context
from local_review_fakes import LocalCase

NEW = ('description_text', 'area_text', 'duration_text', 'organizer_name', 'availability_text',
       'source_schedule_text')
FACTS = dict(area_text='ایرانشهر - سمیه', duration_text='۳ ساعت', organizer_name='سارا نمونه',
             availability_text='تکمیل ظرفیت', source_schedule_text='یکشنبه‌ها ساعت 19 به تاریخ 5،12 و 19 مهر',
             description_text='=توضیح منبع\nخط دوم توضیح', price_text=None)


def result(**fields):
    return canonicalize((context('a', title='دورهمی موسیقی', **{**FACTS, **fields}),))


class SourceFieldOutputTests(LocalCase):
    def setUp(self):
        super().setUp()
        self.result = result()
        self.persist()
        self.document = self.repo.load()
        self.event = self.document['events'][0]

    def test_canonical_resolution_and_sqlite_record(self):
        event = self.result.events[0]
        self.assertEqual(event.status, EventStatus.SOLD_OUT)
        for name in NEW:
            self.assertEqual(self.event[name], FACTS[name])
        self.assertTrue({'area_text', 'description_text', 'source_schedule_text'}
                        <= {p['field'] for p in self.event['field_provenance']})
        self.assertEqual(self.event['status'], 'sold_out')
        self.assertEqual(self.event['summary'], None)
        self.assertEqual(self.document['schema_version'], SCHEMA_VERSION)

    def test_version1_record_payload_loads_with_null_new_fields(self):
        legacy = {key: value for key, value in self.event.items() if key not in NEW}
        event = event_from_record(legacy)
        self.assertTrue(all(getattr(event, name) is None for name in NEW))
        row = dict(zip(EVENT_HEADERS, event_row(event, 'run')))
        self.assertEqual((row['منطقه / محله'], row['توضیحات / معرفی']), ('', ''))

    def test_existing_database_payloads_need_no_ddl_migration(self):
        # Payloads are versioned JSON records: a Stage 9 database written before
        # these fields keeps working, and the SQLite user_version is unchanged.
        connection = sqlite3.connect(self.repo.path)
        try:
            payload = json.loads(connection.execute('SELECT payload FROM event_snapshots').fetchone()[0])
            for name in NEW:
                payload.pop(name)
            with connection:
                connection.execute('UPDATE event_snapshots SET payload=?', (json.dumps(payload),))
            self.assertEqual(connection.execute('PRAGMA user_version').fetchone()[0], DB_VERSION)
        finally:
            connection.close()
        document = self.repo.load()
        self.assertNotIn('area_text', document['events'][0])
        ExcelReviewExporter(self.book, self.repo).export(document)
        CsvReviewExporter(self.root / 'exports').export(document)
        with (self.root / 'exports/latest/events.csv').open(encoding='utf-8-sig', newline='') as handle:
            self.assertEqual(next(csv.DictReader(handle))['area_text'], '')

    def test_xlsx_columns_badge_full_description_and_bounded_rows(self):
        ExcelReviewExporter(self.book, self.repo).export(self.document)
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook['_meta']['B1'].value, WORKBOOK_VERSION)
        sheet = next(s for s in workbook if s.title.startswith('رویدادها'))
        headers = [c.value for c in sheet[1]]
        self.assertEqual(tuple(headers[:len(EVENT_VISIBLE)]), EVENT_VISIBLE)
        self.assertNotIn('خلاصه', headers[:len(EVENT_VISIBLE)])
        self.assertLess(headers.index('مکان'), headers.index('منطقه / محله'))
        self.assertLess(headers.index('منطقه / محله'), headers.index('شهر'))
        value = lambda header: sheet.cell(2, headers.index(header) + 1)
        self.assertEqual(value(DESCRIPTION).value, FACTS['description_text'])
        self.assertEqual(value(DESCRIPTION).data_type, 's')  # formula-looking text stays literal
        self.assertTrue(value(DESCRIPTION).alignment.wrap_text)
        self.assertEqual(value('زمان‌بندی اعلام‌شده').value, FACTS['source_schedule_text'])
        self.assertEqual(value('ظرفیت / وضعیت ثبت‌نام').value, 'تکمیل ظرفیت')
        self.assertEqual(value('ظرفیت / وضعیت ثبت‌نام').fill.fgColor.rgb[-6:], 'FBEED6')
        self.assertEqual(value('وضعیت رویداد').value, 'تکمیل ظرفیت')
        self.assertIsNone(value(DESCRIPTION).hyperlink)
        self.assertLessEqual(sheet.row_dimensions[2].height, ROW_MAX)
        self.assertEqual(sheet.freeze_panes, 'A2')

    def test_description_over_xlsx_limit_is_marked_and_kept_whole_elsewhere(self):
        text = 'متن منبع ' * 5000
        self.result = result(description_text=text)
        self.persist('long', delta=5)
        document = self.repo.load('long')
        ExcelReviewExporter(self.book, self.repo).export(document)
        JsonReviewExporter(self.root / 'exports').export(document)
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = next(s for s in workbook if s.title.startswith('رویدادها'))
        headers = [c.value for c in sheet[1]]
        cell = sheet.cell(2, headers.index(DESCRIPTION) + 1).value
        self.assertLessEqual(len(cell), XLSX_CELL_LIMIT)
        self.assertTrue(cell.endswith(TRUNCATED))
        stored = json.loads((self.root / 'exports/latest/review.json').read_text(encoding='utf-8'))
        self.assertEqual(stored['events'][0]['description_text'], text)

    def test_link_looking_description_is_not_turned_into_a_link(self):
        self.result = result(description_text='https://example.test/more\nجزئیات منبع')
        self.persist('links', delta=5)
        ExcelReviewExporter(self.book, self.repo).export(self.repo.load('links'))
        workbook = load_workbook(self.book)
        self.addCleanup(workbook.close)
        sheet = next(s for s in workbook if s.title.startswith('رویدادها'))
        headers = [c.value for c in sheet[1]]
        cell = sheet.cell(2, headers.index(DESCRIPTION) + 1)
        self.assertEqual((cell.value, cell.hyperlink), ('https://example.test/more\nجزئیات منبع', None))

    def test_previous_workbook_schema_is_upgraded_and_other_versions_refused(self):
        exporter = ExcelReviewExporter(self.book, self.repo)
        exporter.export(self.document)
        workbook = load_workbook(self.book)
        old = next(s for s in workbook if s.title.startswith('رویدادها'))
        before = [list(r) for r in old.values]
        title = old.title
        workbook['_meta']['B1'] = 2
        workbook.save(self.book)
        workbook.close()
        self.persist('two', delta=60)
        exporter.export(self.repo.load('two'))
        workbook = load_workbook(self.book)
        self.assertEqual(workbook['_meta']['B1'].value, WORKBOOK_VERSION)
        self.assertEqual([list(r) for r in workbook[title].values], before)  # history untouched
        for version in (1, WORKBOOK_VERSION + 1):
            workbook['_meta']['B1'] = version
            workbook.save(self.book)
            with self.subTest(version=version), self.assertRaises(ReviewError):
                exporter.export(self.repo.load('two'))
        workbook.close()

    def test_source_fact_edits_are_visual_only(self):
        ExcelReviewExporter(self.book, self.repo).export(self.document)
        workbook = load_workbook(self.book)
        sheet = next(s for s in workbook if s.title.startswith('رویدادها'))
        headers = [c.value for c in sheet[1]]
        for header, value in (('منطقه / محله', 'ویرایش'), (DESCRIPTION, 'ویرایش'), ('مدت', '۹ ساعت'),
                              ('برگزارکننده', 'دیگری'), ('ظرفیت / وضعیت ثبت‌نام', ''),
                              ('تصمیم من', DECISIONS[1])):
            sheet.cell(2, headers.index(header) + 1).value = value
        workbook.save(self.book)
        workbook.close()
        summary = ExcelReviewImporter(self.book, self.repo).import_reviews()
        self.assertEqual((summary.imported, summary.malformed, summary.conflicts), (1, 0, 0))
        event = self.repo.load()['events'][0]
        self.assertEqual(event['review_decision'], DECISIONS[1])
        for name in NEW:
            self.assertEqual(event[name], FACTS[name])

    def test_csv_json_and_gemini_agree_on_every_new_field(self):
        CsvReviewExporter(self.root / 'exports').export(self.document)
        JsonReviewExporter(self.root / 'exports').export(self.document)
        paths = GeminiBundleExporter(self.root / 'exports/gemini').export(self.document)
        with (self.root / 'exports/latest/events.csv').open(encoding='utf-8-sig', newline='') as handle:
            row = next(csv.DictReader(handle))
        stored = json.loads((self.root / 'exports/latest/review.json').read_text(encoding='utf-8'))['events'][0]
        bundle = self.root / 'exports/gemini/latest'
        gemini = json.loads((bundle / 'review.json').read_text(encoding='utf-8'))['events'][0]
        for name in NEW:
            self.assertEqual(stored[name], self.event[name])
            self.assertEqual(gemini[name], self.event[name])
            expected = "'" + self.event[name] if self.event[name].startswith('=') else self.event[name]
            self.assertEqual(row[name], expected)  # formula-risk text is neutralized in CSV only
        manifest = json.loads((bundle / 'manifest.json').read_text(encoding='utf-8'))
        for name, digest in manifest['sha256'].items():
            self.assertEqual(sha256((bundle / name).read_bytes()).hexdigest(), digest)
        self.assertEqual(manifest['export_schema_version'], SCHEMA_VERSION)
        self.assertTrue(any(p.name == 'GEMINI_INSTRUCTIONS.md' for p in paths))

    def test_gemini_instructions_cover_new_semantics(self):
        for phrase in ('description_text', 'untrusted', 'area_text', 'Never build or infer an address',
                       'source_schedule_text', 'never expand', 'availability_text', 'Never invent a capacity',
                       'never generate it'):
            self.assertIn(phrase, INSTRUCTIONS)

    def test_shared_sheets_row_carries_new_fields(self):
        event = replace(self.result.events[0], start_date=None)
        row = dict(zip(EVENT_HEADERS, event_row(event, 'run')))
        self.assertEqual((row['منطقه / محله'], row['مدت'], row['برگزارکننده']),
                         (FACTS['area_text'], FACTS['duration_text'], FACTS['organizer_name']))
        self.assertEqual(row['source_schedule_text'], FACTS['source_schedule_text'])
        self.assertEqual(row['توضیحات / معرفی'], FACTS['description_text'])
        self.assertNotIn('خلاصه', row)


class ScheduleDisplayTests(LocalCase):
    def test_simple_modelled_date_is_not_repeated_but_partial_is_shown(self):
        simple = canonicalize((context('s', source_date_text='۱ مهر ساعت ۱۸'),)).events[0]
        self.assertEqual(dict(zip(EVENT_HEADERS, event_row(simple, 'r')))['زمان‌بندی اعلام‌شده'], '')
        self.assertEqual(dict(zip(EVENT_HEADERS, event_row(simple, 'r')))['source_date_text'], '۱ مهر ساعت ۱۸')
        partial = canonicalize((context('p', source_date_text='جمعه‌ها', start_date=None, start_time=None,
                                        starts_at=None, date_precision=DatePrecision.UNKNOWN),)).events[0]
        self.assertEqual(dict(zip(EVENT_HEADERS, event_row(partial, 'r')))['زمان‌بندی اعلام‌شده'], 'جمعه‌ها')

    def test_listing_detail_conflict_is_a_visible_diagnostic(self):
        event = canonicalize((context('c', field_conflicts=('price_text',)),)).events[0]
        self.assertIn(('source_field_conflict', 'price_text'), {(d.code, d.field) for d in event.diagnostics})
        self.assertEqual(dict(zip(EVENT_HEADERS, event_row(event, 'r')))['کیفیت داده'], 'تعارض داده')

    def test_description_differences_across_sources_are_not_conflicts(self):
        events = canonicalize((context('a', publisher='p', description_text='توضیح یک'),
                               context('b', publisher='p', description_text='توضیح دو'))).events
        event, = events
        self.assertEqual(event.description_text, 'توضیح یک')
        self.assertFalse([d for d in event.diagnostics if d.field == 'description_text'])
