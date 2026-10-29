"""Admission is separate from extraction, scheduling and historical review."""
import csv
import json
from dataclasses import replace
from datetime import date
from unittest import TestCase
from unittest.mock import patch

from openpyxl import load_workbook

from deduplication_fakes import context
from discovery_fakes import make_raw_item
from local_review_fakes import LocalCase, STAMP, sample
from sheets_fakes import RepositoryCase
from test_extract_cli import CONFIG, PUBLISHED_AT, write_store
from test_grouping import fragment
from gatherradar.deduplication import canonicalize
from gatherradar.deduplication.models import CanonicalizationResult
from gatherradar.domain.event import FieldProvenance
from gatherradar.exports.excel import ExcelReviewExporter, ExcelReviewImporter
from gatherradar.exports.gemini import GeminiBundleExporter
from gatherradar.exports.portable import CsvReviewExporter, JsonReviewExporter
from gatherradar.orchestration.canonicalization_run import run_canonical_review
from gatherradar.review.eligibility import eligible, exclusion_reason, eligible_document
from gatherradar.review.records import event_record, dumps
from gatherradar.review.schema import DECISIONS
from gatherradar.review.serialization import ReviewWindow, review_window
from gatherradar.storage import JsonlEvidenceStore
from gatherradar.storage.sqlite_schema import migrate


class TitlePolicyTests(TestCase):
    def test_real_source_title_is_included_unchanged(self):
        event = sample().events[0]
        self.assertTrue(eligible(event))
        self.assertTrue(eligible(event_record(event)))
        self.assertTrue(eligible(replace(event, title='Unknown Event: a theatre performance')))

    def test_blank_placeholder_and_symbol_only_records_are_excluded(self):
        record = event_record(sample().events[0])
        for title in (None, '', ' \t\n', '\u200b\u200f', '---', '🎉',
                      'Untitled', ' UNKNOWN EVENT! ', 'Untitled 🎉', 'رویداد بدون عنوان',
                      'رويداد بدون\u200cعنوان', 'N/A'):
            with self.subTest(title=title):
                self.assertFalse(eligible(dict(record, title=title)))

    def test_source_attribution_is_required_not_model_confidence(self):
        event = sample().events[0]
        for changed in (replace(event, field_provenance=(), extraction_confidence=1),
                        replace(event, source_ids=()),
                        replace(event, field_provenance=(FieldProvenance('title', ('foreign',)),))):
            self.assertEqual(exclusion_reason(changed), 'unsupported_title')


class AdmissionTests(LocalCase):
    def test_persistence_guards_even_a_caller_supplied_unfiltered_window(self):
        good = self.result.events[0]
        missing = replace(self.result.events[1], title=None)
        placeholder = replace(good, event_id='placeholder', title='Untitled')
        result = replace(self.result, events=(good, missing, placeholder))
        snapshot = self.repo.sync(result, ReviewWindow(result.events, 0, 0),
            run_id='one', started_at=STAMP, finished_at=STAMP, days=14,
            review_timezone='Asia/Tehran', source_count=2)
        self.assertEqual(snapshot.events, 1)
        run = self.repo.load()['run']
        self.assertEqual((run['discovered_event_count'], run['excluded_event_count'],
                          run['canonical_event_count'], run['filtered'], run['event_count']),
                         (3, 2, 1, 0, 1))
        self.assertEqual(run['excluded_missing_title_by_source'], {'source_a': 1, 'source_b': 1})
        self.assertEqual(run['title_exclusion_reasons'], {'missing_title': 1, 'placeholder_title': 1})
        for table in ('events', 'event_identity_map', 'event_review', 'event_snapshots'):
            self.assertEqual(self.repo.connection.execute(f'SELECT count(*) FROM {table}').fetchone()[0], 1)
        self.assertEqual(self.repo.load()['possible_duplicates'], [])
        self.assertEqual(result.events[1].venue_name, missing.venue_name)
        self.assertIsNotNone(missing.start_date)  # Date and venue do not replace a title.

    def test_counts_separate_admission_and_date_filtering(self):
        good = self.result.events[0]
        old = replace(self.result.events[1], start_date=date(2025, 1, 1), occurrences=())
        missing = replace(old, event_id='missing', title=None)
        self.persist(result=replace(self.result, events=(good, old, missing)))
        run = self.repo.load()['run']
        self.assertEqual((run['discovered_event_count'], run['excluded_event_count'],
                          run['canonical_event_count'], run['filtered'], run['event_count']),
                         (3, 1, 2, 1, 1))
        self.assertEqual(len(self.repo.load(include_filtered=True)['events']), 2)

    def test_ocr_recovered_titles_and_multiple_events_survive_and_evidence_is_unchanged(self):
        raw_path = write_store(self.root, [('POSTER', PUBLISHED_AT, '')])
        item = make_raw_item('', shortcode='POSTER')
        parts = (fragment(item, 'کارگاه سفالگری\nجمعه ساعت ۱۸', slide=0),
                 fragment(item, 'کنسرت باران\nشنبه ساعت ۲۰', slide=1))
        evidence_path = self.root / 'evidence/instagram/davvvat.jsonl'
        JsonlEvidenceStore(evidence_path).append_new(parts)
        before = (raw_path.read_bytes(), evidence_path.read_bytes())
        review = run_canonical_review(('davvvat_instagram',), config_path=CONFIG,
                                      data_dir=self.root, instagram_evidence=True)
        self.assertEqual(len(review.result.events), 2)
        self.assertTrue(all(eligible(e) for e in review.result.events))
        self.persist(result=review.result)
        self.assertEqual(len(self.repo.load()['schedule_review']), 2)
        self.assertEqual(len(self.repo.load(include_filtered=True)['events']), 2)
        self.assertEqual((raw_path.read_bytes(), evidence_path.read_bytes()), before)
        self.assertEqual(len({e.event_id for e in review.result.events}), 2)

    def legacy(self, title=None):
        self.persist()
        document = self.repo.load()
        ExcelReviewExporter(self.book, self.repo).export(document)
        missing = document['events'][1]
        missing['title'] = title
        # Model a pre-policy SQLite payload and its historical workbook.
        self.repo.connection.execute('UPDATE event_snapshots SET payload=? WHERE event_id=?',
                                     (dumps(missing), missing['event_id']))
        self.repo.connection.execute('UPDATE events SET payload=? WHERE event_id=?',
                                     (dumps(missing), missing['event_id']))
        book = load_workbook(self.book)
        sheet = book.active
        sheet.cell(3, 2, '').value = title
        sheet.cell(3, 1, DECISIONS[1])
        headers = [c.value for c in sheet[1]]
        sheet.cell(3, headers.index('یادداشت من') + 1, 'Keep historical decision')
        book.save(self.book)
        book.close()
        imported = ExcelReviewImporter(self.book, self.repo).import_reviews()
        self.assertEqual(imported.imported, 1)
        return document, missing['event_id']

    def test_legacy_reexport_preserves_database_snapshots_and_reviews(self):
        _, missing_id = self.legacy()
        before = '\n'.join(self.repo.connection.iterdump())
        document = self.repo.load()
        self.assertEqual(len(document['events']), 1)
        self.assertEqual(document['run']['excluded_missing_title_by_source'], {'source_b': 1})
        self.assertEqual(self.repo.find_run('one').events, 1)
        self.assertEqual(self.repo.load()['events'], document['events'])
        self.assertEqual('\n'.join(self.repo.connection.iterdump()), before)
        self.assertEqual(self.repo.connection.execute('SELECT decision FROM event_review WHERE event_id=?',
                                                      (missing_id,)).fetchone()[0], DECISIONS[1])
        migrate(self.repo.connection)
        migrate(self.repo.connection)
        self.assertEqual('\n'.join(self.repo.connection.iterdump()), before)
        self.assertEqual(self.repo.connection.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
        self.assertEqual(self.repo.connection.execute('PRAGMA foreign_key_check').fetchall(), [])

    def test_legacy_whitespace_is_filtered_before_domain_deserialization(self):
        self.legacy(' \t ')
        self.assertEqual(len(self.repo.load()['events']), 1)

    def test_all_exports_share_eligible_ids_and_same_run_xlsx_keeps_history(self):
        legacy, _ = self.legacy()
        before = load_workbook(self.book)
        old_name = before.active.title
        old_values = list(before.active.values)
        before.close()
        # Direct legacy-document exporters must guard too, not rely on load().
        expected = {legacy['events'][0]['event_id']}
        projected = eligible_document(legacy)
        self.assertEqual(eligible_document(projected), projected)
        exports = self.root / 'exports'
        CsvReviewExporter(exports).export(legacy)
        JsonReviewExporter(exports).export(legacy)
        GeminiBundleExporter(exports / 'gemini').export(legacy)
        exporter = ExcelReviewExporter(self.book, self.repo)
        exporter.export(legacy)
        book = load_workbook(self.book)
        self.assertEqual(list(book[old_name].values), old_values)
        self.assertNotEqual(book.active.title, old_name)
        headers = [c.value for c in book.active[1]]
        ids = {row[headers.index('event_id')] for row in book.active.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(ids, expected)
        names = book.sheetnames
        book.close()
        exporter.export(legacy)
        book = load_workbook(self.book)
        self.assertEqual(book.sheetnames, names)
        book.close()
        imported = ExcelReviewImporter(self.book, self.repo).import_reviews()
        self.assertEqual((imported.malformed, imported.conflicts), (0, 0))
        for directory in (exports / 'latest', exports / 'gemini/latest'):
            document = json.loads((directory / 'review.json').read_text(encoding='utf-8'))
            self.assertEqual({e['event_id'] for e in document['events']}, expected)
            with (directory / 'events.csv').open(encoding='utf-8-sig', newline='') as handle:
                self.assertEqual({r['event_id'] for r in csv.DictReader(handle)}, expected)
        manifest = json.loads((exports / 'gemini/latest/manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(manifest['counts']['events'], 1)


class SheetsAdmissionTests(RepositoryCase):
    def test_google_snapshot_and_identity_map_omit_ineligible_events(self):
        from gatherradar.sheets.schema import EVENT_MAP
        good = canonicalize((context(),)).events[0]
        bad = replace(good, event_id='missing', title=None)
        result = CanonicalizationResult(events=(good, bad))
        snapshot = self.repo.sync(result, ReviewWindow(result.events, 0, 0),
            run_id='admission', started_at=STAMP, finished_at=STAMP, days=14,
            review_timezone='Asia/Tehran', source_count=1)
        self.assertEqual(snapshot.events, 1)
        self.assertEqual(len(self.client.values(snapshot.event_tab)), 2)
        self.assertEqual(len(self.client.values(EVENT_MAP)), 2)

    def test_legacy_google_reexport_appends_idempotent_projection_and_keeps_history(self):
        from copy import deepcopy
        good = canonicalize((context(),)).events[0]
        result = CanonicalizationResult(events=(good, replace(good, event_id='missing', title=None)))
        arguments = dict(run_id='legacy', started_at=STAMP, finished_at=STAMP, days=14,
                         review_timezone='Asia/Tehran', source_count=1)
        window = ReviewWindow(result.events, 0, 0)
        with patch('gatherradar.review.serialization.eligible_review', return_value=(result, window)):
            previous = self.repo.sync(result, window, **arguments)
        old = deepcopy(self.client.named(previous.event_tab))
        current = self.repo.sync(result, window, **arguments)
        self.assertEqual(current.events, 1)
        self.assertNotEqual(current.run_id, previous.run_id)
        self.assertEqual(self.client.named(previous.event_tab), old)
        all_sheets = deepcopy(self.client.sheets)
        again = self.repo.sync(result, window, **arguments)
        self.assertEqual(again.run_id, current.run_id)
        self.assertEqual(self.client.sheets, all_sheets)
