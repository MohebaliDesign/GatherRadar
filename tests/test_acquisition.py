"""Publisher → channel → strategy: fallback, provenance, coverage and safety (offline)."""
from __future__ import annotations

import inspect
import os
import unittest
from dataclasses import replace
from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import Mock, patch

from gatherradar.acquisition import coverage as coverage_module, models, policy, strategies
from gatherradar.acquisition.coverage import publisher_coverage, with_channel_gaps
from gatherradar.acquisition.models import ChannelResult, ChannelStatus, classify_failure
from gatherradar.acquisition.policy import recheck_policies
from gatherradar.acquisition.strategies import LinkedPageStrategy, WebsiteListingStrategy, referenced_urls
from gatherradar.collectors.base import (AuthenticationRequiredError, CollectionResult, SourceAccessRestrictedError,
                                         SourceUnavailableError)
from gatherradar.collectors.websites import transport
from gatherradar.config import load_sources
from gatherradar.deduplication.resolution import resolve_group
from gatherradar.domain import SourceType
from gatherradar.domain.website import WebsiteConfig
from gatherradar.extraction import RuleBasedDiscoveryProvider
from gatherradar.extraction.models import ExtractionInput
from gatherradar.orchestration.collection_run import run_instagram_collection
from gatherradar.orchestration.refresh import RefreshService
from gatherradar.orchestration.website_run import run_website_collection
from deduplication_fakes import context
from repo_config import FIXTURE_CONFIG, REPO_CONFIG
from sheets_fakes import STAMP, RepositoryCase
from test_refresh import raw

SHIPPED = {s.id: s for s in load_sources(REPO_CONFIG)}
FIXTURE = {s.id: s for s in load_sources(FIXTURE_CONFIG)}


def instagram_context(key, **fields):
    ctx = context(key, publisher='davvvat', **fields)
    source = replace(ctx.source, source_type=SourceType.INSTAGRAM, username='example')
    return replace(ctx, source=source, raw_item=replace(ctx.raw_item, source_type=SourceType.INSTAGRAM))


def result(source_id, publisher, channel, status, observed=0, strategy=None):
    return ChannelResult(source_id, publisher, channel, strategy or f'{channel}_x', status, observed)


class PublisherRefreshTests(RepositoryCase):
    """Strategy fallback through the real RefreshService with fake collectors."""

    def setUp(self):
        super().setUp()
        self.ig = FIXTURE['davvvat_instagram']
        self.site = FIXTURE['davvvat_website']
        self.evidence = Mock(return_value=SimpleNamespace(failures=(), malformed=()))

    def service(self, *, website=None, instagram=None, **kwargs):
        def collect(source, item_key):
            collector = Mock()
            collector.collect.return_value = CollectionResult(source.id, (raw(source, item_key),))
            return collector
        website = website or (lambda source_id, **kw: run_website_collection(
            source_id, collector=collect(self.site, 'w'), **kw))
        instagram = instagram or (lambda source_id, **kw: run_instagram_collection(
            source_id, collector=collect(self.ig, 'i'), **kw))
        return RefreshService(self.repo, website_collect=website, instagram_collect=instagram,
                              acquire_evidence=self.evidence, now=lambda: STAMP, **kwargs)

    def run_service(self, service, ids, config=FIXTURE_CONFIG, **kwargs):
        return service.run(ids, config_path=config, data_dir=self.directory, run_id=kwargs.pop('run_id', 'pub'),
                           skip_instagram_evidence=True, **kwargs)

    def coverage(self, result):
        return {c.publisher_key: c for c in result.publishers}

    def test_website_success_is_the_normal_flow(self):
        result = self.run_service(self.service(), (self.site.id, self.ig.id))
        self.assertEqual(result.snapshot.status, 'success')
        davvvat = self.coverage(result)['davvvat']
        self.assertTrue(davvvat.covered and davvvat.all_channels_healthy)
        self.assertEqual({(c.strategy, c.status) for c in result.channels},
                         {('website_listing', ChannelStatus.SUCCESS), ('instagram_profile', ChannelStatus.SUCCESS)})

    def test_website_unavailable_instagram_success_covers_publisher(self):
        broken = Mock(side_effect=SourceUnavailableError('down', category='network', operation='request'))
        result = self.run_service(self.service(website=broken), (self.site.id, self.ig.id))
        davvvat = self.coverage(result)['davvvat']
        self.assertTrue(davvvat.covered)
        self.assertFalse(davvvat.all_channels_healthy)
        self.assertEqual(davvvat.covered_by, (self.ig.id,))
        self.assertEqual(result.snapshot.status, 'partial')  # an unexpected failure is still reported
        website = next(c for c in result.channels if c.channel == 'website')
        self.assertEqual(website.status, ChannelStatus.TEMPORARY_FAILURE)
        self.assertIn('description_text', davvvat.potentially_unavailable_fields)

    def test_runtime_robots_block_with_instagram_success_is_not_partial(self):
        blocked = Mock(side_effect=SourceAccessRestrictedError('robots', category='robots', operation='robots'))
        result = self.run_service(self.service(website=blocked), (self.site.id, self.ig.id))
        self.assertEqual(result.snapshot.status, 'success')
        self.assertEqual(next(c for c in result.channels if c.channel == 'website').status,
                         ChannelStatus.POLICY_BLOCKED)
        self.assertTrue(self.coverage(result)['davvvat'].covered)

    def test_configured_policy_block_is_reported_but_never_contacted(self):
        website = Mock()
        result = self.run_service(self.service(website=website), (self.ig.id,), config=REPO_CONFIG)
        website.assert_not_called()
        site = next(c for c in result.channels if c.source_id == 'davvvat_website')
        self.assertEqual((site.status, site.diagnostic, site.observed), (ChannelStatus.POLICY_BLOCKED, 'robots', 0))
        self.assertEqual(result.snapshot.status, 'success')
        record = next(c for c in result.publishers if c.publisher_key == 'davvvat').record()
        self.assertEqual({c['source_id']: c['status'] for c in record['channels']},
                         {'davvvat_instagram': 'success', 'davvvat_website': 'policy_blocked'})

    def test_all_channels_fail_publisher_unavailable(self):
        down = Mock(side_effect=SourceUnavailableError('down', category='network', operation='request'))
        auth = Mock(side_effect=AuthenticationRequiredError('login', category='collection_failure'))
        result = self.run_service(self.service(website=down, instagram=auth), (self.site.id, self.ig.id))
        davvvat = self.coverage(result)['davvvat']
        self.assertFalse(davvvat.covered)
        self.assertEqual(davvvat.potentially_unavailable_fields, ())
        self.assertEqual(result.snapshot.status, 'failed')
        self.assertEqual({c.status for c in result.channels},
                         {ChannelStatus.TEMPORARY_FAILURE, ChannelStatus.ACCESS_RESTRICTED})

    def test_stored_mode_is_offline(self):
        self.run_service(self.service(), (self.site.id, self.ig.id))
        website, instagram, linked, check = Mock(), Mock(), Mock(), Mock()
        service = self.service(website=website, instagram=instagram, linked=linked, policy_check=check)
        result = self.run_service(service, (self.ig.id,), config=REPO_CONFIG, collect=False, run_id='stored')
        for fake in (website, instagram, linked.plan, linked.acquire, check):
            fake.assert_not_called()
        self.assertEqual({c.status for c in result.channels}, {ChannelStatus.STORED, ChannelStatus.POLICY_BLOCKED})
        self.assertTrue(self.coverage(result)['davvvat'].covered)

    def test_policy_recheck_runs_for_disabled_peers_during_collection(self):
        check = Mock(return_value=())
        self.run_service(self.service(policy_check=check), (self.ig.id,), config=REPO_CONFIG)
        (peers,), kwargs = check.call_args
        self.assertEqual([p.id for p in peers], ['davvvat_website'])
        self.assertEqual(kwargs['state_path'], Path(self.directory) / 'state' / 'policy-checks.json')


class ChannelProvenanceTests(unittest.TestCase):
    def test_instagram_only_facts_have_no_website_provenance(self):
        event = resolve_group((instagram_context('a'),), 'group')
        self.assertEqual({(p.channel, p.strategy, p.publisher_key) for p in event.channel_provenance},
                         {('instagram', 'instagram_profile', 'davvvat')})
        venue = next(p for p in event.field_provenance if p.field == 'venue_name')
        channels = {p.candidate_id: p.channel for p in event.channel_provenance}
        self.assertEqual({channels[c] for c in venue.candidate_ids}, {'instagram'})

    def test_conflicting_channels_stay_reviewable(self):
        site = context('w', publisher='davvvat', venue_name='Blue Gallery')
        insta = instagram_context('i', venue_name='Red Hall')
        event = resolve_group((site, insta), 'group')
        self.assertIsNone(event.venue_name)
        self.assertIn(('field_conflict', 'venue_name'), {(d.code, d.field) for d in event.diagnostics})
        self.assertEqual({p.channel for p in event.channel_provenance}, {'website', 'instagram'})

    def test_complementary_channels_keep_their_own_field_origins(self):
        site = context('w', publisher='davvvat', venue_name=None, address='خیابان نمونه ۱')
        insta = instagram_context('i')
        event = resolve_group((site, insta), 'group')
        channels = {p.candidate_id: p.channel for p in event.channel_provenance}
        origin = {p.field: {channels[c] for c in p.candidate_ids} for p in event.field_provenance}
        self.assertEqual(origin['address'], {'website'})
        self.assertEqual(origin['venue_name'], {'instagram'})

    def test_linked_page_strategy_is_recorded(self):
        ctx = context('l', publisher='jabama')
        ctx = replace(ctx, raw_item=replace(ctx.raw_item, raw_metadata={'acquisition_strategy': 'linked_page'}))
        event = resolve_group((ctx,), 'group')
        self.assertEqual(event.channel_provenance[0].strategy, 'linked_page')


class ReferenceUrlTests(unittest.TestCase):
    CAPTION = ('کارگاه سفال مهتاب\n۵ مهر ساعت ۱۸\nمکان: گالری آبی\n'
               'جزئیات: https://publisher.example/event/42 و https://www.instagram.com/p/abc/')

    def discover(self, text):
        return RuleBasedDiscoveryProvider().discover(ExtractionInput(
            'instagram:x:1', 'x', 'instagram', text, 'https://www.instagram.com/p/1/', 'post'))

    def test_instagram_caption_url_is_preserved_for_owner_navigation(self):
        facts = self.discover(self.CAPTION)
        self.assertEqual(facts.reference_urls, ('https://publisher.example/event/42',))  # platform self-link omitted
        insta = instagram_context('i', reference_urls=facts.reference_urls, description_text=None)
        event = resolve_group((insta,), 'group')
        self.assertEqual(event.reference_urls, ('https://publisher.example/event/42',))
        self.assertIsNone(event.description_text)  # the unseen page's content is not claimed
        self.assertNotIn('https://publisher.example/event/42', event.evidence_urls)

    def test_reference_urls_round_trip_through_records(self):
        from gatherradar.review.records import event_from_record, event_record
        event = resolve_group((instagram_context('i', reference_urls=('https://publisher.example/event/42',)),), 'g')
        again = event_from_record(event_record(event))
        self.assertEqual((again.reference_urls, again.channel_provenance), (event.reference_urls, event.channel_provenance))

    def test_records_without_channel_fields_still_load(self):
        from gatherradar.review.records import event_from_record, event_record
        record = event_record(resolve_group((context('a'),), 'g'))
        for name in ('reference_urls', 'channel_provenance', 'channel_gaps'):
            record.pop(name)
        self.assertEqual(event_from_record(record).channel_provenance, ())


class PublisherCoverageTests(unittest.TestCase):
    def test_coverage_is_separate_from_channel_health(self):
        (davvvat,) = publisher_coverage((result('davvvat_website', 'davvvat', 'website', ChannelStatus.UNAVAILABLE),
                                         result('davvvat_instagram', 'davvvat', 'instagram', ChannelStatus.SUCCESS, 5)))
        self.assertTrue(davvvat.covered)
        self.assertFalse(davvvat.all_channels_healthy)
        self.assertEqual(davvvat.unavailable_channels, ('davvvat_website',))

    def test_empty_success_is_healthy_but_not_coverage(self):
        (publisher,) = publisher_coverage((result('a', 'p', 'instagram', ChannelStatus.EMPTY),))
        self.assertFalse(publisher.covered)
        self.assertTrue(publisher.all_channels_healthy)

    def test_same_kind_peer_does_not_reduce_coverage(self):
        (jabama,) = publisher_coverage((result('events', 'jabama', 'website', ChannelStatus.SUCCESS, 5),
                                        result('experiences', 'jabama', 'website', ChannelStatus.PARSE_FAILURE)))
        self.assertEqual(jabama.potentially_unavailable_fields, ())

    def test_channel_gaps_distinguish_unavailable_channel_from_source_omission(self):
        coverage = publisher_coverage((result('davvvat_website', 'davvvat', 'website', ChannelStatus.POLICY_BLOCKED),
                                       result('davvvat_instagram', 'davvvat', 'instagram', ChannelStatus.SUCCESS, 1)))
        event = resolve_group((instagram_context('i', description_text=None, address='آدرس نمونه'),), 'g')
        other = resolve_group((context('o', publisher='other', description_text=None),), 'g2')
        marked, untouched = with_channel_gaps((event, other), coverage)
        self.assertIn('description_text', marked.channel_gaps)
        self.assertNotIn('address', marked.channel_gaps)  # present: never marked
        self.assertNotIn('title', marked.channel_gaps)  # Instagram states titles itself
        self.assertEqual(untouched.channel_gaps, ())  # its publisher lost no channel
        self.assertIsNone(marked.description_text)  # never filled


class FailureVocabularyTests(unittest.TestCase):
    def test_failures_are_not_collapsed(self):
        cases = {
            SourceAccessRestrictedError('r', category='robots', operation='robots'): ChannelStatus.POLICY_BLOCKED,
            SourceAccessRestrictedError('x', category='http_status', http_status=403): ChannelStatus.ACCESS_RESTRICTED,
            SourceAccessRestrictedError('x', category='http_status', http_status=429): ChannelStatus.TEMPORARY_FAILURE,
            AuthenticationRequiredError('login'): ChannelStatus.ACCESS_RESTRICTED,
            SourceUnavailableError('t', category='timeout'): ChannelStatus.TEMPORARY_FAILURE,
            SourceUnavailableError('l', category='layout'): ChannelStatus.PARSE_FAILURE,
            SourceUnavailableError('s', category='http_status', http_status=404): ChannelStatus.UNAVAILABLE,
        }
        for exc, expected in cases.items():
            with self.subTest(exc=exc):
                self.assertEqual(classify_failure(exc), expected)


class NoUnsafeFallbackTests(unittest.TestCase):
    """Architectural guarantees: fallbacks cannot alter identity, proxying or access."""

    def test_transport_identity_and_proxying_are_fixed(self):
        self.assertEqual(transport.USER_AGENT, 'GatherRadar/0.1')
        with patch.dict(os.environ, {'HTTPS_PROXY': 'http://proxy.example:8080', 'HTTP_PROXY': 'http://proxy.example:8080'}):
            http = transport.HttpTransport('https://publisher.example/')
        # ProxyHandler({}) replaces urllib's environment-proxy default: no handler proxies.
        self.assertFalse([h for h in http._opener.handlers if getattr(h, 'proxies', None)])
        self.assertEqual(set(inspect.signature(transport.HttpTransport).parameters), {'source_url'})

    def test_strategies_accept_no_identity_proxy_or_browser_options(self):
        forbidden = {'user_agent', 'headers', 'proxy', 'proxies', 'stealth', 'cookies', 'captcha', 'launcher'}
        for strategy in (strategies.WebsiteListingStrategy, strategies.InstagramProfileStrategy, LinkedPageStrategy):
            for method in (strategy.__init__, strategy.acquire):
                self.assertFalse(forbidden & set(inspect.signature(method).parameters), (strategy, method))

    def test_acquisition_package_imports_no_evasion_tooling(self):
        banned = ('playwright_stealth', 'undetected', 'requests', 'httpx', 'selenium', 'cloudscraper',
                  'fake_useragent', 'socks', 'captcha')
        for module in (models, strategies, coverage_module, policy):
            imports = [line for line in inspect.getsource(module).splitlines()
                       if line.startswith(('import ', 'from '))]
            for line in imports:
                self.assertFalse(any(b in line for b in banned), line)
        self.assertNotIn('User-Agent', inspect.getsource(strategies))

    def test_disabled_channel_is_never_contacted(self):
        collect = Mock()
        outcome = WebsiteListingStrategy(collect).acquire(SHIPPED['davvvat_website'], config_path=REPO_CONFIG,
                                                          data_dir='unused', limit=5)
        collect.assert_not_called()
        self.assertEqual(outcome.status, ChannelStatus.POLICY_BLOCKED)

    def test_linked_page_to_policy_blocked_target_is_not_contacted(self):
        factory = Mock()
        strategy = LinkedPageStrategy(factory)
        outcome = strategy.acquire(SHIPPED['davvvat_website'], {'https://davvvat.ir/event/1': ('x',)},
                                   data_dir='unused', limit=5)
        factory.assert_not_called()
        self.assertEqual(outcome.status, ChannelStatus.POLICY_BLOCKED)


class LinkedPagePlanTests(unittest.TestCase):
    TARGET = replace(FIXTURE['vadoostan_website'])

    def item(self, key, text, source='vadoostan_instagram'):
        return raw(FIXTURE[source], key, text=text)

    def test_one_approved_page_per_observation_and_bounded(self):
        urls = [f'https://vadoostan.ir/app/experiences/Item{i}' for i in range(4)]
        items = [self.item(f'k{i}', f'{urls[i]}\n{urls[(i + 1) % 4]}') for i in range(4)]
        plan = LinkedPageStrategy().plan(items, (self.TARGET,), limit=2)
        self.assertEqual(plan, {'vadoostan_website': {urls[0]: ('vadoostan_instagram:k0',),
                                                      urls[1]: ('vadoostan_instagram:k1',)}})

    def test_unapproved_origins_listings_and_website_seeds_are_ignored(self):
        items = [self.item('a', 'https://evil.example/app/experiences/X'),
                 self.item('b', 'https://vadoostan.ir/app/experiences'),  # listing, not a detail
                 self.item('c', 'https://vadoostan.ir/app/experiences/Y', source='vadoostan_website')]
        self.assertEqual(LinkedPageStrategy().plan(items, (self.TARGET,), limit=5), {})

    def test_referenced_urls_are_verbatim(self):
        item = self.item('a', 'ثبت نام: https://vadoostan.ir/app/experiences/Z، جزئیات')
        self.assertEqual(referenced_urls(item), ('https://vadoostan.ir/app/experiences/Z',))

    def test_linked_collection_uses_the_normal_website_boundary(self):
        from gatherradar.collectors.website import WebsiteCollector
        from test_source_coverage import Pages, V_LIST, html
        source = replace(self.TARGET, url=V_LIST, website=WebsiteConfig('vadoostan', ('/app/experiences/',), 'main'))
        detail = V_LIST + '/GameNightA1'
        collected = WebsiteCollector(transport=Pages({detail: html('vadoostan-experience-detail.html')}),
                                     now=lambda: STAMP).collect_linked(source, (detail,), limit=1,
                                                                       linked_from={detail: ('ig:1',)})
        (item,) = collected.items
        self.assertEqual(item.raw_metadata['acquisition_strategy'], 'linked_page')
        self.assertEqual(item.raw_metadata['linked_from'], ['ig:1'])
        with self.assertRaises(SourceUnavailableError):
            WebsiteCollector(transport=Pages({})).collect_linked(source, ('https://evil.example/x',), limit=1,
                                                                  linked_from={})


class PolicyRecheckTests(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.state = Path(directory.name) / 'policy.json'
        self.sources = list(SHIPPED.values())

    def test_now_allowed_is_reported_never_enabled(self):
        (notice,) = recheck_policies(self.sources, state_path=self.state, now=STAMP, allows=lambda s: True)
        self.assertEqual((notice.source_id, notice.state), ('davvvat_website', 'now_allowed'))
        self.assertFalse({s.id: s for s in load_sources(REPO_CONFIG)}['davvvat_website'].enabled)

    def test_cached_within_interval_and_rechecked_after(self):
        allows = Mock(return_value=False)
        recheck_policies(self.sources, state_path=self.state, now=STAMP, allows=allows)
        (cached,) = recheck_policies(self.sources, state_path=self.state, now=STAMP + timedelta(days=1), allows=allows)
        self.assertEqual((allows.call_count, cached.fresh, cached.state), (1, False, 'still_blocked'))
        recheck_policies(self.sources, state_path=self.state, now=STAMP + timedelta(days=8), allows=allows)
        self.assertEqual(allows.call_count, 2)

    def test_transient_failure_changes_nothing(self):
        recheck_policies(self.sources, state_path=self.state, now=STAMP, allows=lambda s: False)
        before = self.state.read_text(encoding='utf-8')
        def fail(source):
            raise SourceUnavailableError('down', category='network', operation='request')
        (notice,) = recheck_policies(self.sources, state_path=self.state, now=STAMP + timedelta(days=8), allows=fail)
        self.assertEqual((notice.state, notice.detail), ('check_failed', 'request:network'))
        self.assertEqual(self.state.read_text(encoding='utf-8'), before)

    def test_only_robots_disabled_website_channels_are_checked(self):
        allows = Mock(return_value=False)
        enabled = [s for s in self.sources if s.enabled]
        self.assertEqual(recheck_policies(enabled, state_path=self.state, now=STAMP, allows=allows), ())
        allows.assert_not_called()


if __name__ == '__main__':
    unittest.main()


class PublisherExportTests(RepositoryCase):
    """Local workspace outputs keep the channel/provenance distinction."""

    def test_exports_report_channels_and_keep_reference_links(self):
        import csv
        import io
        import json
        from contextlib import redirect_stdout
        from openpyxl import load_workbook
        from gatherradar.review.cli import print_publishers
        from gatherradar.review.workspace import local_workspace
        from gatherradar.storage.sqlite_repository import SqliteCanonicalRepository

        root = Path(self.directory)
        ig = SHIPPED['davvvat_instagram']
        caption = ('کارگاه سفال مهتاب\n۵ مهر ۱۴۰۵ ساعت ۱۸\nمکان: گالری آبی\n'
                   'جزئیات: https://davvvat.ir/event/42')
        item = raw(ig, 'ref', text=caption)
        collector = Mock()
        collector.collect.return_value = CollectionResult(ig.id, (item,))
        with SqliteCanonicalRepository(root / 'state.sqlite3') as repo:
            workspace = local_workspace(repo, root / 'review.xlsx', root / 'exports')
            result = RefreshService(repo, website_collect=Mock(), now=lambda: STAMP,
                instagram_collect=lambda sid, **kw: run_instagram_collection(sid, collector=collector, **kw),
                import_reviews=workspace.import_reviews, export_reviews=workspace.export,
                linked=LinkedPageStrategy(Mock())).run(
                    (ig.id,), config_path=REPO_CONFIG, data_dir=root, run_id='exp', skip_instagram_evidence=True)
        self.assertEqual(result.snapshot.status, 'success')
        linked = next(c for c in result.channels if c.strategy == 'linked_page')
        self.assertEqual(linked.status, ChannelStatus.POLICY_BLOCKED)  # davvvat.ir: reported, not fetched

        document = json.loads((root / 'exports/latest/review.json').read_text(encoding='utf-8'))
        (davvvat,) = [c for c in document['run']['publisher_coverage'] if c['publisher'] == 'davvvat']
        self.assertTrue(davvvat['covered'])
        self.assertIn('davvvat_website', davvvat['unavailable_channels'])
        (event,) = document['events']
        self.assertEqual(event['reference_urls'], ['https://davvvat.ir/event/42'])
        self.assertEqual({p['channel'] for p in event['channel_provenance']}, {'instagram'})
        self.assertIn('description_text', event['channel_gaps'])
        self.assertIsNone(event['description_text'])

        with (root / 'exports/latest/events.csv').open(encoding='utf-8-sig', newline='') as handle:
            (row,) = list(csv.DictReader(handle))
        self.assertIn('davvvat.ir/event/42', row['reference_urls'])
        instructions = next((root / 'exports/gemini/latest').glob('*.md')).read_text(encoding='utf-8')
        self.assertIn('Missing fields caused by unavailable channels must remain missing', instructions)

        book = load_workbook(root / 'review.xlsx')
        links = [c.hyperlink.target for row in book.active.iter_rows() for c in row if c.hyperlink]
        book.close()
        self.assertIn('https://davvvat.ir/event/42', links)  # clickable reference, content not claimed

        output = io.StringIO()
        with redirect_stdout(output):
            print_publishers(result, {'davvvat': 'Davvvat'})
        text = output.getvalue()
        self.assertIn('Website (davvvat_website): policy_blocked', text)
        self.assertIn('Instagram (davvvat_instagram): success — 1 observed', text)
        self.assertIn('publisher_coverage: covered through davvvat_instagram', text)
        self.assertIn('full description', text)
