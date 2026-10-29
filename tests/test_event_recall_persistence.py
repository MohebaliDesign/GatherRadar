"""The source-backed display title contract survives all local projections."""
import csv
import json
from dataclasses import replace

from openpyxl import load_workbook

from gatherradar.deduplication.models import CanonicalizationResult
from gatherradar.review.serialization import review_window
from gatherradar.review.workspace import local_workspace
from local_review_fakes import LocalCase
from test_event_recall import pipeline, REFERENCE


class RecallPersistenceTests(LocalCase):
    def test_derived_title_export_identity_history_and_repeatability(self):
        events, _ = pipeline('در این کارگاه ساخت شمع را یاد می‌گیریم.\n۱ آبان ۱۴۰۵ ساعت ۱۸\nثبت نام')
        self.assertEqual(len(events), 1)
        value = events[0]
        result = CanonicalizationResult(events=events)
        def persist(run_id, result):
            self.repo.sync(result, review_window(result.events, REFERENCE), run_id=run_id,
                started_at=REFERENCE, finished_at=REFERENCE, days=14,
                review_timezone='Asia/Tehran', source_count=1)
        persist('original', result)
        workspace = local_workspace(self.repo, self.book, self.root / 'exports')
        self.assertFalse(workspace.export('original')[1])
        original = load_workbook(self.book)
        old_name, old_rows = original.active.title, list(original.active.values)
        original.close()
        persist('replay', replace(result, events=(replace(value, description_text=value.description_text + '\nتوضیح'),)))
        self.assertFalse(workspace.export('replay')[1])
        document = self.repo.load('replay')
        row = document['events'][0]
        self.assertEqual(row['event_id'], value.event_id)
        self.assertIsNone(row['source_title'])
        self.assertEqual(row['display_title_origin'], 'derived_from_source_description')
        self.assertEqual(set(row['display_title_evidence_ids']), set(value.display_title_evidence_ids))
        for folder in (self.root / 'exports/latest', self.root / 'exports/gemini/latest'):
            self.assertEqual(json.loads((folder / 'review.json').read_text(encoding='utf-8')), document)
            with (folder / 'events.csv').open(encoding='utf-8-sig', newline='') as stream:
                self.assertEqual([r['event_id'] for r in csv.DictReader(stream)], [value.event_id])
        book = load_workbook(self.book)
        self.assertNotEqual(book.active.title, old_name)
        self.assertEqual(list(book[old_name].values), old_rows)
        self.assertTrue(book.active.sheet_view.rightToLeft)
        headers = [c.value for c in book.active[1]]
        self.assertEqual(book.active.cell(2, headers.index('event_id') + 1).value, value.event_id)
        names = book.sheetnames
        book.close()
        self.assertFalse(workspace.export('replay')[1])
        book = load_workbook(self.book)
        self.assertEqual(book.sheetnames, names)
        book.close()
        self.assertEqual(self.repo.load('original')['events'][0]['description_text'], value.description_text)
        self.assertEqual(self.repo.connection.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
