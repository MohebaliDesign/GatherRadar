from dataclasses import replace
from urllib.error import URLError
import socket
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from gatherradar.exports.excel import ExcelReviewExporter, ExcelReviewImporter, open_workspace, RUNS, DUPLICATES
from gatherradar.exports.excel_design import column_style, font_name, ROW_MIN, ROW_MAX
from gatherradar.exports.portable import portable_files
from gatherradar.review.models import ReviewError
from gatherradar.collectors.base import SourceUnavailableError, safe_failure
from gatherradar.collectors.websites.transport import HttpTransport
from gatherradar.extraction import DiscoveryService, RuleBasedDiscoveryProvider
from gatherradar.extraction.service import build_extraction_input
from gatherradar.extraction.fields import source_date_text, address, venue_name
from gatherradar.extraction.signals import analyze
from test_website import collect, FakeTransport, BASE, SOURCE
from local_review_fakes import LocalCase


class DesignTests(LocalCase):
    def setUp(self):
        super().setUp()
        self.persist()
        self.exporter=ExcelReviewExporter(self.book,self.repo)
        self.exporter.export(self.repo.load())

    def read(self):
        book=load_workbook(self.book)
        self.addCleanup(book.close)
        return book

    def test_fonts_header_density_and_widths(self):
        for sheet in self.read():
            if sheet.title=='_meta': continue
            self.assertEqual(sheet.freeze_panes,'A2')
            self.assertTrue(sheet.sheet_view.rightToLeft)
            self.assertFalse(sheet.merged_cells.ranges)
            self.assertTrue(26<=sheet.row_dimensions[1].height<=32)
            for cell in sheet[1]:
                self.assertTrue(cell.font.bold)
                self.assertNotEqual(cell.fill.fgColor.rgb,sheet.cell(2,cell.column).fill.fgColor.rgb)
                spec=column_style(cell.value)
                self.assertTrue(spec.minimum<=sheet.column_dimensions[cell.column_letter].width<=spec.maximum)
            for row in sheet.iter_rows(min_row=2):
                self.assertTrue(ROW_MIN<=sheet.row_dimensions[row[0].row].height<=ROW_MAX)
                for cell in row:
                    self.assertEqual(cell.font.name,font_name(cell.value))

    def test_latest_navigation_history_and_old_export(self):
        old=self.read().active.title
        self.persist('two',delta=60)
        self.exporter.export(self.repo.load('two'))
        book=self.read()
        newest=book.active.title
        self.assertNotEqual(newest,old)
        self.assertEqual(book.sheetnames[0],newest)
        self.assertEqual(book.sheetnames[2:4],[RUNS,DUPLICATES])
        self.assertIn(old,book.sheetnames)
        self.assertEqual(book['_meta'].sheet_state,'hidden')
        self.assertTrue(all(len(n)<=31 for n in book.sheetnames))
        self.exporter.export(self.repo.load('one'))
        self.assertEqual(self.read().active.title,newest)

    def test_badges_defaults_and_technical_columns(self):
        book=self.read()
        sheet=book.active
        headers=[c.value for c in sheet[1]]
        self.assertGreater(len(sheet.conditional_formatting),0)
        self.assertTrue(sheet.data_validations.dataValidation)
        self.assertEqual(sheet['A2'].value,'بررسی نشده')
        self.assertIsNone(sheet.cell(2,headers.index('یادداشت من')+1).value)
        for header in ('event_id','category','status','event_format','registration_url','review_token'):
            self.assertTrue(sheet.column_dimensions[get_column_letter(headers.index(header)+1)].hidden)
        self.assertEqual(font_name('English فارسی'),'Vazir')
        self.assertEqual(font_name('event:abc'),'Poppins')

    def test_long_url_compact_label_and_canonical_exports_unchanged(self):
        document=self.repo.load()
        target='https://example.test/'+'x'*2000
        document['events'][0]['registration_url']=target
        before=portable_files(document)
        path=self.root/'long.xlsx'
        ExcelReviewExporter(path,self.repo).export(document)
        book=load_workbook(path)
        self.addCleanup(book.close)
        sheet=book.active
        header=[c.value for c in sheet[1]]
        cell=sheet.cell(2,header.index('ثبت‌نام / خرید')+1)
        self.assertEqual(cell.value,'ثبت‌نام / خرید')
        self.assertEqual(cell.hyperlink.target,target)
        self.assertLessEqual(sheet.column_dimensions[cell.column_letter].width,20)
        self.assertEqual(before,portable_files(document))

    def test_schema_one_refused_without_overwrite(self):
        book=self.read()
        book['_meta']['B1']=1
        book.save(self.book)
        before=self.book.read_bytes()
        with self.assertRaises(ReviewError): open_workspace(self.book)
        with self.assertRaises(ReviewError): ExcelReviewImporter(self.book,self.repo).import_reviews()
        self.assertEqual(before,self.book.read_bytes())


class RecallTests(unittest.TestCase):
    def test_explicit_heading_survives_without_relaxing_classification(self):
        raw=collect().items[0]
        text='اجرای آبی\nکنسرت روز ۱۲ مهر ساعت ۱۸ برگزار می‌شود.'
        raw=replace(raw,raw_text=text,raw_metadata={'page_title':'اجرای آبی','title_origin':'heading'})
        outcome=DiscoveryService(RuleBasedDiscoveryProvider()).discover(raw,SOURCE)
        self.assertEqual(outcome.event.title,'اجرای آبی')
        unrelated=replace(raw,raw_text='اجرای آبی\nاین صرفاً نام نوشته است.')
        self.assertIsNone(DiscoveryService(RuleBasedDiscoveryProvider()).discover(unrelated,SOURCE).event)

    def test_fallback_and_non_evidence_heading_not_passed(self):
        raw=collect().items[0]
        for metadata in ({'page_title':'fake','title_origin':'heading'},
                         {'page_title':raw.raw_text.splitlines()[0],'title_origin':'fallback'},
                         {'page_title':['invalid'],'title_origin':'heading'}):
            self.assertIsNone(build_extraction_input(replace(raw,raw_metadata=metadata)).source_title)

    def test_adapter_heading_is_explicit_and_affects_observation_hash(self):
        raw=collect().items[0]
        self.assertEqual(raw.raw_metadata['title_origin'],'heading')
        page='<article id="event"><p>کارگاه نقاشی آبی</p><p>۱۲ مهر ساعت ۱۸</p></article>'
        fallback=collect(FakeTransport({BASE:'<a href="/events/one">One</a>',BASE+'events/one':page})).items[0]
        self.assertEqual(fallback.raw_metadata['title_origin'],'fallback')
        self.assertIsNone(build_extraction_input(fallback).source_title)

    def test_standalone_address_section_preserved_not_prose(self):
        self.assertEqual(address(analyze('نمایشگاه آبی\nآدرس\nتهران، خیابان نمونه')),'تهران، خیابان نمونه')
        self.assertIsNone(address(analyze('برای آدرس با ما تماس بگیرید\nتهران')))

    def test_recurrence_qualifier_does_not_drop_clock(self):
        self.assertEqual(source_date_text(analyze('شنبه ها، ساعت 21')),'شنبه ها، ساعت 21')

    def test_directional_location_is_address_not_venue_name(self):
        found=analyze('مکان: روبروی پارک نمونه، بازارچه آبی')
        self.assertIsNone(venue_name(found))
        self.assertEqual(address(found),'روبروی پارک نمونه، بازارچه آبی')


class SafeDiagnosticsTests(unittest.TestCase):
    def test_exception_payload_never_logged(self):
        self.assertEqual(safe_failure(RuntimeError('secret=https://private/token')), 'collect:collection_failure')
        self.assertEqual(safe_failure(SourceUnavailableError('private',category='http_status',operation='robots',http_status=403)), 'robots:http_status:http_403')
        self.assertEqual(safe_failure(SourceUnavailableError('private',category='SECRET',operation='SECRET',http_status='secret')), 'collect:collection_failure')

    def test_network_categories(self):
        for cause,category in ((socket.gaierror('private'),'dns'),(TimeoutError('private'),'timeout'),(OSError('private'),'network')):
            transport=HttpTransport('https://example.test/')
            with patch.object(transport._opener,'open',side_effect=URLError(cause)):
                with self.assertRaises(SourceUnavailableError) as caught:
                    transport._request('https://example.test/')
            self.assertEqual(safe_failure(caught.exception),'request:'+category)
