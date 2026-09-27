import unittest
from dataclasses import replace
from datetime import date, datetime, time, timezone
from pathlib import Path

from persiantools.jdatetime import JalaliDate
from openpyxl import load_workbook

from gatherradar.collectors.website import WebsiteCollector
from gatherradar.collectors.websites.transport import Page
from gatherradar.domain import Source, SourceType, RawItem, EventCandidate
from gatherradar.domain.website import WebsiteConfig
from gatherradar.domain.evidence import primary_discovery_units
from gatherradar.domain.temporal import DatePrecision
from gatherradar.extraction import DiscoveryService, RuleBasedDiscoveryProvider
from gatherradar.extraction.fields import venue_name, address, city, event_format
from gatherradar.extraction.signals import analyze
from gatherradar.normalization import normalize_temporal, NormalizationService
from gatherradar.normalization.ordinals import DAY_ORDINALS
from gatherradar.deduplication import CandidateContext, canonicalize
from gatherradar.exports.excel import ExcelReviewExporter
from gatherradar.review.serialization import review_window
from local_review_fakes import LocalCase, STAMP

SOURCE = Source('sample', 'publisher', 'Sample', SourceType.WEBSITE,
                'https://example.test/', website=WebsiteConfig(content_selector='.event-detail-right', detail_path_prefixes=('/events/',)))
RAW = RawItem('raw:sample',SOURCE.id,SOURCE.source_type,'sample','webpage',
              'https://example.test/events/sample','',datetime(2026,9,26,12,tzinfo=timezone.utc))
COMBINED = 'یکم تا سوم مهرماه از ساعت ۱۰ تا ۲۲'


def temporal(text, raw=RAW):
    return normalize_temporal(text,raw,SOURCE)


def codes(result):
    return {d.code for d in result.diagnostics}


class PersianDateWordsTests(unittest.TestCase):
    def test_required_ranges_and_prefix(self):
        for text in ('یکم تا سوم مهرماه','اول تا سوم مهر','۱ تا ۳ مهر','از یکم تا سوم مهر'):
            with self.subTest(text=text):
                result=temporal(text)
                self.assertEqual((result.start_date,result.end_date),(date(2026,9,23),date(2026,9,25)))
                self.assertEqual(result.date_precision,DatePrecision.RANGE)
                self.assertIn('inferred_year',codes(result))

    def test_all_ordinals_through_valid_days(self):
        for name,day in DAY_ORDINALS.items():
            with self.subTest(name=name):
                self.assertEqual(temporal(name+' فروردین ۱۴۰۵').start_date,JalaliDate(1405,1,day).to_gregorian())

    def test_month_variants_all_months(self):
        for number,month in enumerate('فروردین اردیبهشت خرداد تیر مرداد شهریور مهر آبان آذر دی بهمن اسفند'.split(),1):
            for suffix in ('','ماه',' ماه'):
                self.assertEqual(temporal('اول '+month+suffix+' ۱۴۰۵').start_date,JalaliDate(1405,number,1).to_gregorian())

    def test_orthographic_variants(self):
        for name,day in (('بیست‌ویکم',21),('بیست و دوم',22),('سی‌ام',30),('سی‌اُم',30),('سوّم',3)):
            self.assertEqual(temporal(name+' شهریور ۱۴۰۵').start_date,JalaliDate(1405,6,day).to_gregorian())

    def test_explicit_year_and_numeric_regression(self):
        for text in ('بیست و پنجم تا بیست و هفتم شهریور ۱۴۰۵','۲۵ تا ۲۷ شهریور ۱۴۰۵'):
            result=temporal(text)
            self.assertEqual((result.start_date,result.end_date),(date(2026,9,16),date(2026,9,18)))
            self.assertNotIn('inferred_year',codes(result))

    def test_published_reference_wins_over_capture(self):
        raw=replace(RAW,published_at=datetime(2025,9,24,tzinfo=timezone.utc))
        result=temporal(COMBINED,raw)
        self.assertEqual(result.start_date,JalaliDate(1404,7,1).to_gregorian())
        self.assertEqual(result.reference_basis,'published_at')

    def test_capture_reference_when_unpublished(self):
        self.assertEqual(temporal(COMBINED).reference_basis,'captured_at')

    def test_inference_stays_bounded(self):
        result=temporal('اول تا سوم فروردین')
        self.assertIsNone(result.start_date)
        self.assertIn('missing_year',codes(result))

    def test_discrete_days_are_not_continuous_range(self):
        for text in ('۱ و ۲ مهر','اول و دوم مهرماه','یکم، دوم مهر','بیست و یکم و بیست و دوم مهر'):
            result=temporal(text)
            self.assertIsNone(result.start_date)
            self.assertIn('multiple_dates',codes(result))

    def test_invalid_days_never_partially_match(self):
        for text in ('سی و یکم مهر ۱۴۰۵','سی و دوم مهر ۱۴۰۵','صفرم مهر','۲یکم مهر','دوازدهمین مهر','اول مهرماهانه'):
            with self.subTest(text=text):
                self.assertIsNone(temporal(text).start_date)
                self.assertTrue(temporal(text).diagnostics)

    def test_esfand_calendar_validation(self):
        self.assertEqual(temporal('سی ام اسفند ۱۳۹۹').start_date,date(2021,3,20))
        self.assertIsNone(temporal('سی ام اسفند ۱۴۰۰').start_date)
        self.assertIn('invalid_date',codes(temporal('سی ام اسفند ۱۴۰۰')))

    def test_cross_nowruz_ambiguous_unresolved(self):
        raw=replace(RAW,captured_at=datetime(2026,3,20,tzinfo=timezone.utc))
        result=temporal('بیست و نهم اسفند تا دوم فروردین',raw)
        self.assertIsNone(result.start_date)
        self.assertIn('missing_year',codes(result))

    def test_clear_time_ranges(self):
        for wording,end in (('از ساعت ۱۰ تا ۲۲',time(22)),('ساعت ۱۰ تا ۲۲',time(22))):
            result=temporal(wording)
            self.assertEqual((result.start_time,result.end_time),(time(10),end))
        result=temporal('از ساعت ۱۰:۳۰ تا ۲۲:۱۵')
        self.assertEqual((result.start_time,result.end_time),(time(10,30),time(22,15)))
        result=temporal('از ساعت ۱۰:۳۰ تا ۲۲')
        self.assertEqual((result.start_time,result.end_time),(time(10,30),time(22)))

    def test_invalid_time_ranges(self):
        for wording in ('از ساعت ۲۵ تا ۲۲','از ساعت ۱۰:۷۰ تا ۲۲:۱۵'):
            result=temporal(wording)
            self.assertIsNone(result.start_time)
            self.assertIn('invalid_time',codes(result))

    def test_combined_components_never_continuous_multiday_interval(self):
        result=temporal(COMBINED)
        self.assertEqual((result.start_date,result.end_date),(date(2026,9,23),date(2026,9,25)))
        self.assertEqual((result.start_time,result.end_time),(time(10),time(22)))
        self.assertIsNone(result.starts_at)
        self.assertIsNone(result.ends_at)
        self.assertEqual(result.date_precision,DatePrecision.RANGE)
        self.assertEqual(codes(result),{'inferred_year','multi_day_hours'})

    def test_recurring_negative(self):
        for wording in ('شنبه ها، ساعت 21','شنبه‌ها، ساعت ۲۱','دوشنبه تا جمعه هر هفته'):
            result=temporal(wording)
            self.assertIsNone(result.start_date)
            self.assertIsNone(result.starts_at)
            self.assertTrue(result.diagnostics)


class VenueAndFormatTests(unittest.TestCase):
    def test_explicit_venue_and_address_both_retained(self):
        found=analyze('مکان: خانه آبی\nآدرس: تهران، خیابان نمونه، پلاک ۹')
        self.assertEqual(venue_name(found),'خانه آبی')
        self.assertEqual(address(found),'تهران، خیابان نمونه، پلاک ۹')
        self.assertEqual(city(found),'تهران')

    def test_address_only_never_uses_final_token_as_venue(self):
        found=analyze('آدرس: تهران، خیابان نمونه، خانه آبی')
        self.assertIsNone(venue_name(found))

    def test_venue_only(self):
        found=analyze('محل برگزاری: خانه آبی')
        self.assertEqual(venue_name(found),'خانه آبی')
        self.assertIsNone(address(found))

    def test_prose_alone_or_organizer_alone_insufficient(self):
        for text in ('برنامه اکران در خانه آبی، فرصتی برای دیدار است.\nآدرس: خیابان نمونه، خانه آبی',
                     'برگزارکننده\nخانه آبی\nآدرس: خیابان نمونه، خانه آبی'):
            self.assertIsNone(venue_name(analyze(text)))

    def test_ambiguous_location_address_stays_address(self):
        found=analyze('مکان: تهران، خیابان نمونه، خانه آبی')
        self.assertIsNone(venue_name(found))
        self.assertEqual(address(found),'تهران، خیابان نمونه، خانه آبی')

    def test_organizer_plus_occurrence_and_address_corroboration(self):
        text='برگزارکننده\nخانه آبی\nبرنامه اکران در خانه آبی، فرصتی برای دیدار است.\nآدرس: خیابان نمونه، خانه آبی'
        self.assertEqual(venue_name(analyze(text)),'خانه آبی')
        self.assertIsNone(venue_name(analyze(text.replace('برنامه اکران','برنامه قبلی'))))

    def test_ticket_purchase_and_registration_do_not_set_online(self):
        for text in ('خرید بلیط آنلاین','ثبت نام آنلاین','فروشگاه آنلاین','آدرس: خیابان نمونه\nخرید بلیط آنلاین'):
            self.assertIsNone(event_format(analyze(text)))

    def test_direct_online_participation(self):
        for text in ('به صورت آنلاین برگزار می‌شود','شرکت آنلاین','فرمت: مجازی','attend online','online workshop'):
            self.assertEqual(event_format(analyze(text)),'online')

    def test_transaction_format_is_not_attendance_format(self):
        for text in ('ثبت نام به صورت آنلاین', 'پرداخت به صورت حضوری است',
                     'خرید بلیط به شکل آنلاین خواهد بود'):
            self.assertIsNone(event_format(analyze(text)))

    def test_historical_and_negated_formats(self):
        for text in ('کارگاه قبلی به صورت آنلاین برگزار شد','این جلسه آنلاین نیست','previous online workshop'):
            self.assertIsNone(event_format(analyze(text)))

    def test_hybrid_needs_both_direct_participation_modes(self):
        self.assertEqual(event_format(analyze('حضوری\nخرید بلیط آنلاین')),'in_person')
        self.assertEqual(event_format(analyze('شرکت آنلاین\nبه صورت حضوری')),'hybrid')
        self.assertIsNone(event_format(analyze('آدرس: خیابان نمونه\nداستان فضای مجازی')))

    def test_joint_hybrid_attendance_but_not_joint_registration(self):
        for text in ('کارگاه حضوری و آنلاین', 'به صورت آنلاین و حضوری برگزار می‌شود',
                     'hybrid workshop', 'فرمت: hybrid'):
            self.assertEqual(event_format(analyze(text)), 'hybrid')
        self.assertIsNone(event_format(analyze('ثبت نام حضوری و آنلاین')))


class AccuracyPipelineTests(LocalCase):
    def test_venue_fixture_through_adapter_units_candidates_storage_and_xlsx(self):
        html=(Path(__file__).parent/'fixtures/websites/venue-address-detail.html').read_text(encoding='utf8')
        class Transport:
            def fetch(self,url,**kwargs):
                return Page(url,'<a href="/events/sample">Event</a>' if url==SOURCE.url else html)
        raw=WebsiteCollector(transport=Transport(),now=lambda:RAW.captured_at).collect(SOURCE,limit=1).items[0]
        units=primary_discovery_units(raw)
        discovery=DiscoveryService(RuleBasedDiscoveryProvider()).discover_unit(raw,units[0],SOURCE)
        self.assertEqual(discovery.event.venue_name,'خانه آبی')
        self.assertIsNone(discovery.event.event_format)
        normalized=NormalizationService().normalize(discovery.event,raw,SOURCE)
        result=canonicalize((CandidateContext(normalized,discovery.event,raw,SOURCE,raw.captured_at,'primary'),))
        self.persist(result=result)
        doc=self.repo.load()
        self.assertEqual(doc['events'][0]['venue_name'],'خانه آبی')
        self.assertIn('خیابان نمونه',doc['events'][0]['address'])
        ExcelReviewExporter(self.book,self.repo).export(doc)
        book=load_workbook(self.book)
        self.addCleanup(book.close)
        headers=[c.value for c in book.active[1]]
        self.assertEqual(book.active.cell(2,headers.index('مکان')+1).value,'خانه آبی')
        self.assertIn('خیابان نمونه',book.active.cell(2,headers.index('آدرس')+1).value)

    def test_expired_art_center_range_filtered_but_historically_persisted(self):
        candidate=EventCandidate('candidate:range',RAW.id,True,title='ایونت آبی',source_date_text=COMBINED,evidence_url=RAW.content_url)
        normalized=NormalizationService().normalize(candidate,RAW,SOURCE)
        self.assertEqual(normalized.candidate.source_date_text,COMBINED)
        result=canonicalize((CandidateContext(normalized,candidate,RAW,SOURCE,RAW.captured_at,'primary'),))
        window=review_window(result.events,STAMP,14)
        self.assertEqual((len(window.events),window.filtered,window.undated),(0,1,0))
        self.persist(result=result)
        self.assertEqual(self.repo.load()['events'],[])
        event=self.repo.load(include_filtered=True)['events'][0]
        self.assertEqual((event['start_date'],event['end_date']),('2026-09-23','2026-09-25'))
        self.assertEqual((event['start_time'],event['end_time']),('10:00:00','22:00:00'))
        self.assertEqual(event['source_date_text'],COMBINED)
        self.assertIsNone(event['starts_at'])
        self.assertIsNone(event['ends_at'])
