"""Frozen schedule boundaries and complete persistence/export membership checks."""
import csv
import json
from dataclasses import replace
from datetime import date, datetime, time, timedelta, timezone
from unittest import TestCase
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from persiantools.jdatetime import JalaliDate

from local_review_fakes import LocalCase, sample
from gatherradar.domain.event import CanonicalDiagnostic, EventStatus
from gatherradar.domain.occurrence import EventOccurrence, ScheduleKind
from gatherradar.domain.temporal import DatePrecision
from gatherradar.exports.excel import ExcelReviewExporter, ExcelReviewImporter
from gatherradar.exports.gemini import GeminiBundleExporter
from gatherradar.exports.portable import CsvReviewExporter, JsonReviewExporter
from gatherradar.review.relevance import assess_relevance, relevant_document
from gatherradar.review.schema import DECISIONS
from gatherradar.review.serialization import review_window

ZONE = ZoneInfo('Asia/Tehran')
NOW = datetime(2026, 10, 22, 15, 30, tzinfo=ZONE)  # 30 Mehr 1405
TODAY = NOW.date()


def event(key='sample', **changes):
    base = replace(sample().events[0], event_id=key, start_date=TODAY,
                   end_date=None, start_time=None, end_time=None, starts_at=None, ends_at=None,
                   occurrences=(), schedule_kind=ScheduleKind.SINGLE, recurrence_rule=None,
                   timezone='Asia/Tehran', diagnostics=(), date_precision=DatePrecision.DAY)
    return replace(base, **changes)


def session(key, offset=0, **changes):
    return EventOccurrence(key, TODAY + timedelta(days=offset), 'Asia/Tehran', **changes)


def frozen_events():
    return (
        event('yesterday', start_date=TODAY-timedelta(days=1)),
        event('earlier_today', start_time=time(9), end_time=time(11)),
        event('later_today', start_time=time(18), end_time=time(20)),
        event('ongoing', start_date=TODAY-timedelta(days=3), end_date=TODAY+timedelta(days=16),
              date_precision=DatePrecision.RANGE, start_time=time(8), end_time=time(14)),
        event('ends_today', start_date=TODAY-timedelta(days=3), end_date=TODAY,
              date_precision=DatePrecision.RANGE),
        event('future', start_date=TODAY+timedelta(days=2)),
        event('distant', start_date=TODAY+timedelta(days=15)),
        event('mixed', start_date=None, date_precision=DatePrecision.UNKNOWN,
              schedule_kind=ScheduleKind.DISCRETE,
              occurrences=(session('old', -3), session('next', 1), session('far', 20))),
        event('recurring', start_date=None, date_precision=DatePrecision.UNKNOWN,
              schedule_kind=ScheduleKind.RECURRING, recurrence_rule='FREQ=WEEKLY;BYDAY=TH'),
        event('undated', start_date=None, date_precision=DatePrecision.UNKNOWN),
        event('unknown_end', start_time=time(8)),
        event('cancelled', status=EventStatus.CANCELLED),
        event('date_only'),
    )


class RuntimeBoundaryTests(TestCase):
    def assertDisposition(self, value, expected, reference=NOW, days=14):
        actual = assess_relevance(value, reference, days)
        self.assertEqual(actual.disposition, expected, actual)
        return actual

    def test_frozen_partition_is_exclusive_and_complete(self):
        values = frozen_events()
        window = review_window(values, NOW)
        from gatherradar.review.relevance import disposition_counts
        self.assertEqual(disposition_counts(window.dispositions), dict(actionable=7, expired=2,
            beyond_horizon=1, needs_schedule_review=2, otherwise_excluded=1))
        self.assertEqual(len(window.dispositions), len(values))
        self.assertEqual(len({d.event_id for d in window.dispositions}), len(values))
        self.assertEqual((len(window.events), window.filtered, window.undated), (7, 6, 2))

    def test_end_equal_to_refresh_is_expired(self):
        self.assertDisposition(event(start_time=time(14), end_time=time(15, 30)), 'expired')

    def test_ends_a_second_after_refresh_is_retained(self):
        self.assertDisposition(event(start_time=time(14), end_time=time(15, 30, 1)), 'actionable')

    def test_final_range_closing_time_controls_expiry(self):
        value = event(start_date=TODAY-timedelta(days=5), end_date=TODAY,
                      date_precision=DatePrecision.RANGE, start_time=time(8), end_time=time(14))
        self.assertDisposition(value, 'expired')
        self.assertDisposition(replace(value, end_time=time(18)), 'actionable')

    def test_ongoing_range_ignores_daily_closed_hours_before_final_day(self):
        value = event(start_date=TODAY-timedelta(days=3), end_date=TODAY+timedelta(days=30),
                      date_precision=DatePrecision.RANGE, start_time=time(8), end_time=time(10))
        self.assertDisposition(value, 'actionable')
        self.assertIsNone(value.ends_at)

    def test_date_only_today_and_unknown_end_today_remain_actionable(self):
        for value in (event(), event(start_time=time(7))):
            self.assertDisposition(value, 'actionable')
            self.assertIsNone(value.ends_at)
        self.assertDisposition(event(start_date=TODAY-timedelta(days=1), start_time=time(7)), 'expired')

    def test_horizon_is_inclusive_by_tehran_day(self):
        self.assertDisposition(event(start_date=TODAY+timedelta(days=14)), 'actionable')
        self.assertDisposition(event(start_date=TODAY+timedelta(days=15)), 'beyond_horizon')

    def test_verified_occurrences_outrank_expired_parent(self):
        value = event(start_date=TODAY-timedelta(days=5),
            occurrences=(session('old', -2), session('future', 1), session('far', 25)))
        result = self.assertDisposition(value, 'actionable')
        self.assertEqual(result.occurrence_ids, ('future',))
        self.assertEqual([r[1] for r in result.occurrence_dispositions],
                         ['expired', 'actionable', 'beyond_horizon'])

    def test_exact_session_end_and_unknown_session_end(self):
        start = NOW.replace(hour=9)
        ended = session('ended', starts_at=start, ends_at=start.replace(hour=10),
                        start_time=time(9, 30), end_time=time(10, 30), date_precision=DatePrecision.EXACT)
        self.assertDisposition(event(occurrences=(ended,)), 'expired')
        self.assertDisposition(event(occurrences=(replace(ended, ends_at=None, end_time=None),)), 'actionable')

    def test_unresolved_occurrence_does_not_hide_verified_future_occurrence(self):
        result = self.assertDisposition(event(occurrences=(session('bad', diagnostics=('occurrence_temporal_conflict',)),
            session('good', 1))), 'actionable')
        self.assertEqual(result.occurrence_ids, ('good',))

    def test_price_diagnostic_and_sold_out_are_not_schedule_expiry(self):
        self.assertDisposition(event(status=EventStatus.SOLD_OUT, availability_text='sold out',
            occurrences=(session('sold', 1, remaining_seats=0, enabled=False,
                                 diagnostics=('embedded_session_discount_unresolved',)),)), 'actionable')

    def test_recurring_weekday_has_no_synthetic_sessions(self):
        value = event(start_date=None, date_precision=DatePrecision.UNKNOWN,
                      schedule_kind=ScheduleKind.RECURRING, recurrence_rule='FREQ=WEEKLY;BYDAY=TH')
        self.assertDisposition(value, 'needs_schedule_review')
        self.assertEqual(value.occurrences, ())

    def test_temporal_conflict_unknown_date_and_ambiguous_clock_require_review(self):
        for value in (event(start_date=None), event(schedule_kind=ScheduleKind.CONFLICT),
            event(diagnostics=(CanonicalDiagnostic('normalization:unparsed_temporal_text', 'temporal'),)),
            event(start_time=time(23), end_time=time(2))):
            self.assertDisposition(value, 'needs_schedule_review')

    def test_cancellation_and_postponement_override_future_sessions(self):
        for status in (EventStatus.CANCELLED, EventStatus.POSTPONED):
            self.assertDisposition(event(status=status, occurrences=(session('future', 1),)), 'otherwise_excluded')

    def test_exact_zone_conversion_and_machine_timezone_independence(self):
        reference = datetime(2026, 10, 22, 21, tzinfo=timezone.utc)  # 23 Oct, 00:30 Tehran
        self.assertDisposition(event(), 'expired', reference)
        same_instant = reference.astimezone(ZoneInfo('America/New_York'))
        self.assertEqual(assess_relevance(event(), reference), assess_relevance(event(), same_instant))
        self.assertEqual(assess_relevance(event(), reference, review_timezone='UTC').disposition, 'expired')

    def test_other_event_timezone_is_preserved_at_tehran_boundary(self):
        reference = datetime(2026, 10, 22, 21, tzinfo=timezone.utc)
        self.assertDisposition(event(timezone='America/New_York'), 'actionable', reference)

    def test_persian_year_boundary_and_leap_day(self):
        last = JalaliDate(1403, 12, 30).to_gregorian()
        new_year = JalaliDate(1404, 1, 1).to_gregorian()
        reference = datetime.combine(new_year, time(0, 30), ZONE)
        self.assertDisposition(event(start_date=last), 'expired', reference)
        self.assertDisposition(event(start_date=last, end_date=new_year,
                                    date_precision=DatePrecision.RANGE), 'actionable', reference)

    def test_invalid_reference_and_horizon_rejected(self):
        for reference, days in ((NOW.replace(tzinfo=None), 14), (NOW, 0), (NOW, True)):
            with self.assertRaises(ValueError):
                assess_relevance(event(), reference, days)


class RuntimeExportTests(LocalCase):
    def save(self, run='frozen', reference=NOW):
        result = replace(self.result, events=frozen_events(), places=(), possible_duplicates=())
        return self.repo.sync(result, review_window(result.events, reference), run_id=run,
            started_at=reference, finished_at=reference, days=14, review_timezone='Asia/Tehran', source_count=1)

    def test_sqlite_xlsx_csv_json_gemini_and_history_agree(self):
        self.save()
        document = self.repo.load()
        self.assertEqual(len(document['events']), 7)
        self.assertEqual(len(document['schedule_review']), 2)
        self.assertEqual(len(self.repo.load(include_filtered=True)['events']), 13)
        self.assertEqual(document, relevant_document(document))
        rows = self.repo.connection.execute('SELECT visible,count(*) FROM event_snapshots GROUP BY visible').fetchall()
        self.assertEqual(dict(rows), {0: 6, 1: 7})
        original = '\n'.join(self.repo.connection.iterdump())
        self.repo.load()
        self.assertEqual('\n'.join(self.repo.connection.iterdump()), original)
        exports = self.root / 'exports'
        CsvReviewExporter(exports).export(document)
        JsonReviewExporter(exports).export(document)
        GeminiBundleExporter(exports / 'gemini').export(document)
        exporter = ExcelReviewExporter(self.book, self.repo)
        exporter.export(document)
        expected = {r['event_id'] for r in document['events']}
        for directory in (exports / 'latest', exports / 'gemini/latest'):
            actual = json.loads((directory / 'review.json').read_text(encoding='utf-8'))
            self.assertEqual(actual, document)
            self.assertEqual({r['event_id'] for r in actual['events']}, expected)
            with (directory / 'events.csv').open(encoding='utf-8-sig', newline='') as handle:
                self.assertEqual({r['event_id'] for r in csv.DictReader(handle)}, expected)
            with (directory / 'schedule_review.csv').open(encoding='utf-8-sig', newline='') as handle:
                self.assertEqual({r['event_id'] for r in csv.DictReader(handle)}, {'undated', 'recurring'})
        mixed = next(e for e in document['events'] if e['event_id'] == 'mixed')
        self.assertEqual([o['occurrence_id'] for o in mixed['occurrences']], ['next'])
        stored = next(e for e in self.repo.load(include_filtered=True)['events'] if e['event_id'] == 'mixed')
        self.assertEqual(len(stored['occurrences']), 3)
        book = load_workbook(self.book)
        sheet = book.active
        headers = [c.value for c in sheet[1]]
        self.assertEqual({r[headers.index('event_id')] for r in sheet.iter_rows(min_row=2, values_only=True)}, expected)
        history = list(sheet.values)
        history_name = sheet.title
        review = next(s for s in book if s.title.startswith('زمان‌بندی'))
        self.assertEqual(review.max_row, 3)
        review.cell(2, 1, DECISIONS[1])
        note = [c.value for c in review[1]].index('یادداشت من') + 1
        review.cell(2, note, 'investigate date')
        book.save(self.book)
        names = book.sheetnames
        book.close()
        self.assertEqual(ExcelReviewImporter(self.book, self.repo).import_reviews().imported, 1)
        self.assertIn('investigate date', [r['review_notes'] for r in self.repo.load()['schedule_review']])
        exporter.export(self.repo.load())
        book = load_workbook(self.book)
        self.assertEqual(book.sheetnames, names)
        book.close()
        self.save('new-refresh', NOW + timedelta(days=4))
        self.assertNotEqual(self.repo.load()['run']['event_count'], 7)
        exporter.export(self.repo.load())
        book = load_workbook(self.book)
        self.assertEqual(list(book[history_name].values), history)
        book.close()
        self.assertEqual(self.repo.load('frozen')['run']['event_count'], 7)
        self.assertEqual(self.repo.connection.execute('PRAGMA integrity_check').fetchone()[0], 'ok')
        self.assertEqual(self.repo.connection.execute('PRAGMA foreign_key_check').fetchall(), [])

    def test_legacy_run_projection_uses_stored_clock_and_keeps_records(self):
        self.save()
        # Simulate a pre-policy run with all title-eligible rows in its view.
        self.repo.connection.execute('UPDATE event_snapshots SET visible=1')
        before = '\n'.join(self.repo.connection.iterdump())
        first = self.repo.load('frozen')
        self.assertEqual(first['run']['event_count'], 7)
        self.assertEqual(self.repo.load('frozen'), first)
        self.assertEqual('\n'.join(self.repo.connection.iterdump()), before)

    def test_title_guard_cannot_be_bypassed_by_a_previous_disposition_ledger(self):
        self.save()
        document = self.repo.load()
        document['events'][0]['title'] = None
        document['schedule_review'][0]['title'] = 'Untitled'
        guarded = relevant_document(document)
        self.assertEqual(guarded['run']['event_count'], len(guarded['events']))
        self.assertEqual(guarded['run']['canonical_event_count'], 11)
        self.assertEqual(guarded['run']['excluded_missing_title'], 2)
        self.assertEqual(len(guarded['schedule_review']), 1)
        self.assertEqual(relevant_document(guarded), guarded)
